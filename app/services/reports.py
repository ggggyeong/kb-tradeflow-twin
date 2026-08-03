from __future__ import annotations

import hashlib
import html
import json
import re
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    HRFlowable,
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import PROJECT_ROOT
from app.db.models import (
    DailyMonitoringReport,
    Document,
    FinancialEvent,
    MonitoringRun,
    PaymentObligation,
    Report,
    Shipment,
    TradeCase,
    TransactionFinancialEventLink,
)
from app.services.financial_exposure import FinancialExposureService
from app.services.financial_reference import read_transaction_financial_reference
from app.services.product_advisory import (
    product_availability_label,
    product_requirement_labels,
)

REPORT_DIR = PROJECT_ROOT / "data" / "reports"
BRIEFING_WORKBOOK_PATH = PROJECT_ROOT / "kb_doc" / "KB_TradeFlow_Twin_브리핑_구성안.xlsx"
FONT_PATH = Path("/System/Library/Fonts/Supplemental/AppleGothic.ttf")
KB_LOGO_PATH = PROJECT_ROOT / "app" / "ui" / "assets" / "KB_SymbolMark.png"

KB_NAVY = "#08233F"
KB_BLUE = "#1266B3"
KB_YELLOW = "#FFB900"
KB_DEEP_YELLOW = "#E9A400"
KB_RED = "#B83B45"
KB_ORANGE = "#D97706"
KB_PALE_BLUE = "#EAF3FB"
KB_PALE_YELLOW = "#FFF4CF"
KB_PALE_RED = "#FDEBEC"
KB_PALE_GRAY = "#F4F6F8"
KB_LINE = "#D8E0E8"


def _stable_id(prefix: str, value: str) -> str:
    return f"{prefix}-{hashlib.sha256(value.encode()).hexdigest()[:16]}"


def _hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _register_font() -> str:
    if "NotoSansGothic" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("NotoSansGothic", str(FONT_PATH)))
    return "NotoSansGothic"


def _sheet_sections(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    sheet = workbook[sheet_name]
    sections: list[dict[str, Any]] = []
    for row in sheet.iter_rows(values_only=True):
        order_text = str(row[0] or "").strip()
        if not order_text.isdigit() or not row[1]:
            continue
        sections.append(
            {
                "order": int(order_text),
                "block": str(row[1]).strip(),
                "rules": str(row[2] or "").strip(),
            }
        )
    return sorted(sections, key=lambda item: item["order"])


def load_briefing_blueprint(
    path: Path = BRIEFING_WORKBOOK_PATH,
) -> dict[str, Any]:
    """Load the reviewed eight-block customer/RM workbook contract."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    common: list[dict[str, str]] = []
    for row in workbook["1.공통설계원칙"].iter_rows(values_only=True):
        order_text = str(row[0] or "").strip()
        if not order_text.isdigit() or not row[1]:
            continue
        common.append(
            {
                "principle": str(row[1]).strip(),
                "detail": str(row[2] or "").strip(),
            }
        )
    return {
        "source_path": str(path),
        "source_sha256": _hash_file(path),
        "customer_sections": _sheet_sections(workbook, "2.고객용_구성안"),
        "rm_sections": _sheet_sections(workbook, "3.KB직원용_구성안"),
        "common_principles": common,
    }


def _footer(canvas: Any, document: Any) -> None:
    canvas.saveState()
    canvas.setFont("NotoSansGothic", 8)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawString(20 * mm, 12 * mm, "KB TradeFlow Twin · 상담용 브리핑")
    canvas.drawRightString(190 * mm, 12 * mm, f"{document.page}")
    canvas.restoreState()


def _daily_monitoring_footer(canvas: Any, document: Any) -> None:
    canvas.saveState()
    canvas.setFont("NotoSansGothic", 8)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawString(20 * mm, 12 * mm, "KB TradeFlow Twin · 일일 모니터링")
    canvas.drawRightString(190 * mm, 12 * mm, f"{document.page}")
    canvas.restoreState()


def _safe(value: Any) -> str:
    if value in (None, ""):
        return "-"
    return html.escape(str(value))


def _payload_hash(payload: dict[str, Any]) -> str:
    encoded = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        default=str,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _table(
    rows: list[list[Any]],
    *,
    font: str,
    widths: list[float],
    header_color: str = "#0B1F3A",
) -> Table:
    header_style = ParagraphStyle(
        "TableHeader",
        fontName=font,
        fontSize=7.5,
        leading=9.5,
        textColor=colors.white,
    )
    cell_style = ParagraphStyle(
        "TableCell",
        fontName=font,
        fontSize=7.2,
        leading=9.5,
        textColor=colors.HexColor("#243447"),
        splitLongWords=True,
    )
    rendered_rows = [
        [
            (
                cell
                if isinstance(cell, Paragraph)
                else Paragraph(str(cell), header_style if row_index == 0 else cell_style)
            )
            for cell in row
        ]
        for row_index, row in enumerate(rows)
    ]
    table = Table(rendered_rows, colWidths=widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(header_color)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C5D1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def _display_date(value: Any) -> str:
    """Return an ISO date as a customer-readable Korean date."""
    if value in (None, ""):
        return "미확정"
    raw = str(value)
    try:
        parsed = date.fromisoformat(raw[:10])
    except ValueError:
        return raw
    return f"{parsed.year}년 {parsed.month}월 {parsed.day}일"


def _display_amount(amount: Any, currency: Any) -> str:
    if amount in (None, ""):
        return "금액 확인 필요"
    try:
        number = f"{float(amount):,.0f}"
    except (TypeError, ValueError):
        number = str(amount)
    return f"{number} {str(currency or '').strip()}".strip()


def _first_conflict(payload: dict[str, Any]) -> dict[str, Any]:
    risk = payload.get("risk_snapshot") or {}
    conflicts = [item for item in risk.get("conflicts", []) if isinstance(item, dict)]
    return conflicts[0] if conflicts else {}


def _document_inventory(payload: dict[str, Any]) -> dict[str, dict[str, Any]]:
    inventory: dict[str, dict[str, Any]] = {}
    for item in payload.get("documents", []):
        if not isinstance(item, dict):
            continue
        inventory[str(item.get("doc_type") or "UNKNOWN")] = item
    return inventory


def _briefing_styles(font: str) -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "KBBriefingTitle",
            parent=base["Title"],
            fontName=font,
            fontSize=23,
            leading=30,
            alignment=TA_LEFT,
            textColor=colors.HexColor(KB_NAVY),
            spaceAfter=2 * mm,
        ),
        "subtitle": ParagraphStyle(
            "KBBriefingSubtitle",
            parent=base["BodyText"],
            fontName=font,
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#667587"),
        ),
        "heading": ParagraphStyle(
            "KBBriefingHeading",
            parent=base["Heading2"],
            fontName=font,
            fontSize=15,
            leading=21,
            textColor=colors.HexColor(KB_NAVY),
            spaceBefore=5 * mm,
            spaceAfter=2.5 * mm,
            keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "KBBriefingBody",
            parent=base["BodyText"],
            fontName=font,
            fontSize=9.5,
            leading=15,
            textColor=colors.HexColor("#243447"),
        ),
        "body_large": ParagraphStyle(
            "KBBriefingBodyLarge",
            parent=base["BodyText"],
            fontName=font,
            fontSize=10,
            leading=14.5,
            textColor=colors.HexColor("#172B3F"),
        ),
        "small": ParagraphStyle(
            "KBBriefingSmall",
            parent=base["BodyText"],
            fontName=font,
            fontSize=7.4,
            leading=10.5,
            textColor=colors.HexColor("#667587"),
        ),
        "table_header": ParagraphStyle(
            "KBBriefingTableHeader",
            fontName=font,
            fontSize=8.2,
            leading=11,
            alignment=TA_CENTER,
            textColor=colors.HexColor(KB_NAVY),
        ),
        "table_cell": ParagraphStyle(
            "KBBriefingTableCell",
            fontName=font,
            fontSize=8.4,
            leading=12.5,
            textColor=colors.HexColor("#243447"),
        ),
        "metric": ParagraphStyle(
            "KBBriefingMetric",
            fontName=font,
            fontSize=16,
            leading=21,
            alignment=TA_CENTER,
            textColor=colors.HexColor(KB_NAVY),
        ),
        "metric_label": ParagraphStyle(
            "KBBriefingMetricLabel",
            fontName=font,
            fontSize=7.5,
            leading=10,
            alignment=TA_CENTER,
            textColor=colors.HexColor(KB_NAVY),
        ),
        "callout_title": ParagraphStyle(
            "KBBriefingCalloutTitle",
            fontName=font,
            fontSize=12,
            leading=17,
            textColor=colors.HexColor(KB_NAVY),
            spaceAfter=2 * mm,
        ),
    }


def _briefing_table(
    rows: list[list[Any]],
    *,
    styles: dict[str, ParagraphStyle],
    widths: list[float],
    header_color: str = KB_YELLOW,
) -> Table:
    rendered: list[list[Any]] = []
    for row_index, row in enumerate(rows):
        rendered.append(
            [
                cell
                if isinstance(cell, (Paragraph, Table, Image))
                else Paragraph(
                    str(cell),
                    styles["table_header"] if row_index == 0 else styles["table_cell"],
                )
                for cell in row
            ]
        )
    table = Table(rendered, colWidths=widths, repeatRows=1, hAlign="LEFT")
    commands: list[tuple[Any, ...]] = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(header_color)),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor(KB_LINE)),
        ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor(KB_LINE)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]
    for index in range(1, len(rows)):
        commands.append(
            (
                "BACKGROUND",
                (0, index),
                (-1, index),
                colors.white if index % 2 else colors.HexColor(KB_PALE_GRAY),
            )
        )
    table.setStyle(TableStyle(commands))
    return table


def _briefing_callout(
    title: str,
    body: str,
    *,
    styles: dict[str, ParagraphStyle],
    background: str = KB_PALE_YELLOW,
    accent: str = KB_DEEP_YELLOW,
) -> Table:
    table = Table(
        [
            [Paragraph(title, styles["callout_title"])],
            [Paragraph(body, styles["body_large"])],
        ],
        colWidths=[176 * mm],
        hAlign="LEFT",
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(background)),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor(accent)),
                ("LINEBEFORE", (0, 0), (0, -1), 4, colors.HexColor(accent)),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return table


def _metric_cards(
    items: list[tuple[str, str]],
    *,
    styles: dict[str, ParagraphStyle],
    color: str,
) -> Table:
    cells: list[Any] = []
    width = 176 * mm / max(1, len(items))
    for value, label in items:
        cells.append(
            Table(
                [
                    [Paragraph(_safe(value), styles["metric"])],
                    [Paragraph(_safe(label), styles["metric_label"])],
                ],
                colWidths=[width],
            )
        )
    table = Table([cells], colWidths=[width] * len(items), hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(color)),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.white),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.white),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return table


def _kb_briefing_page(canvas: Any, document: Any, audience: str) -> None:
    canvas.saveState()
    canvas.setFillColor(colors.HexColor(KB_YELLOW))
    canvas.rect(0, A4[1] - 18 * mm, A4[0], 18 * mm, stroke=0, fill=1)
    canvas.setFillColor(colors.HexColor(KB_NAVY))
    canvas.rect(0, A4[1] - 19.2 * mm, A4[0], 1.2 * mm, stroke=0, fill=1)
    canvas.setFillColor(colors.HexColor(KB_NAVY))
    canvas.roundRect(
        15 * mm,
        A4[1] - 16.5 * mm,
        23 * mm,
        14 * mm,
        3 * mm,
        stroke=0,
        fill=1,
    )
    if KB_LOGO_PATH.is_file():
        canvas.drawImage(
            str(KB_LOGO_PATH),
            17 * mm,
            A4[1] - 15 * mm,
            width=19 * mm,
            height=12 * mm,
            preserveAspectRatio=True,
            anchor="c",
            mask="auto",
        )
    canvas.setFont("NotoSansGothic", 9.5)
    canvas.setFillColor(colors.HexColor(KB_NAVY))
    canvas.drawString(40 * mm, A4[1] - 11.2 * mm, "KB TradeFlow Twin")
    label = "고객용 거래 위험 안내" if audience == "CUSTOMER" else "KB 직원용 상담 검토"
    canvas.drawRightString(A4[0] - 17 * mm, A4[1] - 11.2 * mm, label)
    canvas.setStrokeColor(colors.HexColor(KB_LINE))
    canvas.line(17 * mm, 14 * mm, A4[0] - 17 * mm, 14 * mm)
    canvas.setFont("NotoSansGothic", 7.5)
    canvas.setFillColor(colors.HexColor(KB_NAVY))
    footer = "고객 안내용" if audience == "CUSTOMER" else "KB 내부 상담 준비용"
    canvas.drawString(17 * mm, 9 * mm, f"KB TradeFlow Twin · {footer}")
    canvas.drawRightString(A4[0] - 17 * mm, 9 * mm, str(document.page))
    canvas.restoreState()


def _humanize_anchor(value: Any) -> str:
    labels = {
        "ON_BOARD_DATE": "본선적재일 (On Board Date)",
        "BL_DATE": "B/L 발급일",
        "INVOICE_DATE": "Invoice 발행일",
        "FIXED_DATE": "확정일",
    }
    raw = str(value or "").strip().upper()
    return labels.get(raw, raw or "확인 필요")


def _humanize_day_type(value: Any) -> str:
    labels = {
        "CALENDAR": "달력일 (Calendar Day)",
        "BUSINESS": "영업일 (Business Day)",
    }
    raw = str(value or "").strip().upper()
    return labels.get(raw, raw or "확인 필요")


def _doc_status_rows(payload: dict[str, Any]) -> list[list[str]]:
    inventory = _document_inventory(payload)
    return [
        [
            "Booking Confirmation",
            "보유" if "BOOKING_CONFIRMATION" in inventory else "미보유",
        ],
        [
            "Commercial Invoice",
            "보유" if "COMMERCIAL_INVOICE" in inventory else "미보유",
        ],
        [
            "B/L (Bill of Lading)",
            "보유" if "BILL_OF_LADING" in inventory else "미수령",
        ],
    ]


def _risk_numbers(payload: dict[str, Any]) -> dict[str, Any]:
    risk = payload.get("risk_snapshot") or {}
    conflict = _first_conflict(payload)
    frontier = conflict.get("frontier_days_remaining")
    try:
        exceeded_days = max(0, -int(frontier))
    except (TypeError, ValueError):
        exceeded_days = 0
    return {
        "risk": risk,
        "conflict": conflict,
        "planned_due": conflict.get("planned_due_date"),
        "revised_due": conflict.get("revised_due_date"),
        "event_date": conflict.get("event_date"),
        "latest_safe_anchor": conflict.get("latest_safe_anchor_date"),
        "exceeded_days": exceeded_days,
        "priority": (
            conflict.get("priority")
            or conflict.get("response_priority_level")
            or risk.get("highest_priority")
            or "미분류"
        ),
    }


def _product_story(
    payload: dict[str, Any],
    *,
    styles: dict[str, ParagraphStyle],
    rm: bool,
) -> list[Any]:
    options = [item for item in payload.get("product_options", []) if isinstance(item, dict)]
    if not options:
        return [
            _briefing_callout(
                "상품 후보는 상담 단계에서 확인합니다",
                (
                    "대출 만기 조정, 운전자금 확보, 수출채권 활용 가능성을 함께 검토해 주세요. "
                    "실제 이용 가능 여부와 한도·금리·승인은 KB 상담 및 심사 후 확정됩니다."
                ),
                styles=styles,
                background=KB_PALE_YELLOW,
                accent=KB_YELLOW,
            )
        ]
    blocks: list[Any] = []
    for option in options:
        citations = [item for item in option.get("citations", []) if isinstance(item, dict)]
        source_text = ", ".join(
            f"{item.get('source_file', '원문 PDF')} p.{item.get('page', '-')}"
            for item in citations
        ) or "원문 페이지 확인 필요"
        availability = option.get("availability_label") or product_availability_label(
            str(option.get("availability") or "")
        )
        excerpts = ""
        if rm:
            excerpt_lines: list[str] = []
            for item in citations:
                raw_excerpt = str(item.get("excerpt") or "").strip()
                if not raw_excerpt:
                    continue
                shortened = (
                    f"{raw_excerpt[:240].rstrip()}…"
                    if len(raw_excerpt) > 240
                    else raw_excerpt
                )
                excerpt_lines.append(f"원문: {_safe(shortened)}")
            excerpts = "<br/>".join(
                excerpt_lines
            )
        body = (
            f"<b>{_safe(option.get('canonical_name'))}</b><br/>"
            f"{_safe(option.get('why_consider'))}<br/>"
            f"검토 가능 시점: {_safe(availability)}<br/>"
            f"근거: {_safe(source_text)}"
        )
        if excerpts:
            body += f"<br/>{excerpts}"
        blocks.extend(
            [
                KeepTogether(
                    [
                        _briefing_callout(
                            "KB 상품 검토 후보",
                            body,
                            styles=styles,
                            background=KB_PALE_YELLOW,
                            accent=KB_DEEP_YELLOW,
                        )
                    ]
                ),
                Spacer(1, 2 * mm),
            ]
        )
    return blocks


def _priority_sort_key(value: str | None) -> tuple[int, int, str]:
    normalized = str(value or "").strip().upper()
    if normalized.startswith("P") and normalized[1:].isdigit():
        return (0, int(normalized[1:]), normalized)
    if normalized:
        return (1, 0, normalized)
    return (2, 0, "UNRANKED")


def _daily_case_details(monitoring_payload: dict[str, Any]) -> list[dict[str, Any]]:
    """Merge graph ranking with its full per-case evidence without recalculation."""
    ranked = [
        dict(item)
        for item in monitoring_payload.get("ranked_risks", [])
        if isinstance(item, dict) and item.get("case_id")
    ]
    per_case = monitoring_payload.get("per_case_results", {})
    if not isinstance(per_case, dict):
        per_case = {}
    if not per_case:
        task_results = monitoring_payload.get("results", {})
        if isinstance(task_results, dict):
            collected: dict[str, dict[str, Any]] = {}
            for task_id, result in task_results.items():
                if not isinstance(result, dict) or not result.get("case_id"):
                    continue
                case_id = str(result["case_id"])
                entry = collected.setdefault(case_id, {})
                if result.get("source_kind") == "MONITORING_SCAN":
                    entry["financial_risk_snapshot"] = result
                elif str(task_id).startswith("shipment_snapshot_"):
                    entry["shipment_snapshot"] = result
            per_case = collected

    if not ranked:
        for case_id, result in per_case.items():
            if not isinstance(result, dict):
                continue
            snapshot = result.get("financial_risk_snapshot")
            if not isinstance(snapshot, dict):
                continue
            ranked.append(
                {
                    "case_id": str(case_id),
                    "calculation_id": snapshot.get("calculation_id"),
                    "highest_priority": snapshot.get("highest_priority"),
                    "conflict_count": snapshot.get("conflict_count", 0),
                    "conflicts": snapshot.get("conflicts", []),
                }
            )
        ranked.sort(
            key=lambda item: (
                _priority_sort_key(
                    str(item["highest_priority"]) if item.get("highest_priority") else None
                ),
                -int(item.get("conflict_count") or 0),
                str(item["case_id"]),
            )
        )

    details: list[dict[str, Any]] = []
    for position, item in enumerate(ranked, start=1):
        case_id = str(item["case_id"])
        evidence = per_case.get(case_id, {})
        if not isinstance(evidence, dict):
            evidence = {}
        snapshot = evidence.get("financial_risk_snapshot")
        if not isinstance(snapshot, dict):
            snapshot = {}
        shipment = evidence.get("shipment_snapshot")
        if not isinstance(shipment, dict):
            shipment = {}
        conflicts = snapshot.get("conflicts", item.get("conflicts", []))
        if not isinstance(conflicts, list):
            conflicts = []
        data_actions = snapshot.get("data_actions", item.get("data_actions", []))
        if not isinstance(data_actions, list):
            data_actions = []
        details.append(
            {
                "rank": int(item.get("rank") or position),
                "case_id": case_id,
                "calculation_id": (snapshot.get("calculation_id") or item.get("calculation_id")),
                "input_fingerprint": snapshot.get("input_fingerprint"),
                "basis_version": snapshot.get("basis_version"),
                "highest_priority": (
                    snapshot.get("highest_priority") or item.get("highest_priority")
                ),
                "conflict_count": len(conflicts),
                "conflicts": conflicts,
                "data_actions": data_actions,
                "shipment_snapshot": shipment,
            }
        )
    return details


def _render_daily_monitoring_pdf(path: Path, payload: dict[str, Any]) -> None:
    """Render an internal portfolio report from one frozen monitoring payload."""
    font = _register_font()
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "DailyMonitoringTitle",
        parent=styles["Title"],
        fontName=font,
        fontSize=22,
        leading=28,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#0B1F3A"),
        spaceAfter=7 * mm,
    )
    heading = ParagraphStyle(
        "DailyMonitoringHeading",
        parent=styles["Heading2"],
        fontName=font,
        fontSize=13,
        leading=18,
        textColor=colors.HexColor("#006D5B"),
        spaceBefore=5 * mm,
        spaceAfter=2 * mm,
        keepWithNext=True,
    )
    body = ParagraphStyle(
        "DailyMonitoringBody",
        parent=styles["BodyText"],
        fontName=font,
        fontSize=9,
        leading=14,
        textColor=colors.HexColor("#243447"),
    )
    small = ParagraphStyle(
        "DailyMonitoringSmall",
        parent=body,
        fontSize=7.5,
        leading=11,
        textColor=colors.HexColor("#64748B"),
    )
    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=17 * mm,
        leftMargin=17 * mm,
        topMargin=17 * mm,
        bottomMargin=20 * mm,
        title=f"KB TradeFlow Twin Daily Monitoring {payload['as_of_date']}",
        author="KB TradeFlow Twin",
    )
    summary = payload["summary"]
    story: list[Any] = [
        Paragraph("KB TradeFlow Twin", title),
        Paragraph("일일 모니터링 리포트", heading),
        Paragraph(
            (
                f"기준일 {_safe(payload['as_of_date'])} · "
                f"Monitoring Run {_safe(payload['monitoring_run_id'])}"
            ),
            small,
        ),
        Spacer(1, 3 * mm),
        Paragraph(_safe(summary["summary_text"]), body),
        Paragraph("포트폴리오 요약", heading),
        _table(
            [
                ["후보 거래", "선정 거래", "분석 결과", "신규 Alert", "위험 거래", "최고 우선순위"],
                [
                    _safe(summary["candidate_count"]),
                    _safe(summary["selected_count"]),
                    _safe(summary["result_count"]),
                    _safe(summary["new_alert_count"]),
                    _safe(summary["risk_case_count"]),
                    _safe(summary.get("highest_priority")),
                ],
            ],
            font=font,
            widths=[27 * mm, 27 * mm, 27 * mm, 27 * mm, 27 * mm, 27 * mm],
            header_color="#006D5B",
        ),
    ]

    priority_counts = summary.get("priority_summary", {}).get("by_highest_priority", {})
    if isinstance(priority_counts, dict) and priority_counts:
        story.extend(
            [
                Paragraph("우선순위 분포", heading),
                _table(
                    [
                        ["우선순위", "거래 수"],
                        *[
                            [_safe(priority), _safe(count)]
                            for priority, count in priority_counts.items()
                        ],
                    ],
                    font=font,
                    widths=[80 * mm, 81 * mm],
                ),
            ]
        )

    ranked_cases = payload.get("ranked_cases", [])
    story.append(Paragraph("위험 거래 순위", heading))
    if ranked_cases:
        story.append(
            _table(
                [
                    ["순위", "Case", "최고 우선순위", "충돌", "확인 필요", "선적 상태"],
                    *[
                        [
                            _safe(item["rank"]),
                            _safe(item["case_id"]),
                            _safe(item.get("highest_priority")),
                            _safe(item["conflict_count"]),
                            _safe(len(item.get("data_actions", []))),
                            _safe(item.get("shipment_snapshot", {}).get("status")),
                        ]
                        for item in ranked_cases
                    ],
                ],
                font=font,
                widths=[18 * mm, 39 * mm, 28 * mm, 21 * mm, 25 * mm, 31 * mm],
            )
        )
    else:
        story.append(Paragraph("표시할 위험 거래가 없습니다.", body))

    for item in ranked_cases:
        story.append(
            Paragraph(
                f"{_safe(item['rank'])}. {_safe(item['case_id'])}",
                heading,
            )
        )
        story.append(
            Paragraph(
                (
                    f"calculation_id={_safe(item.get('calculation_id'))} · "
                    f"basis={_safe(item.get('basis_version'))} · "
                    f"input_fingerprint={_safe(item.get('input_fingerprint'))}"
                ),
                small,
            )
        )
        conflicts = item.get("conflicts", [])
        if conflicts:
            story.append(
                _table(
                    [
                        ["우선순위", "금융 이벤트", "날짜/간격", "판정 근거"],
                        *[
                            [
                                _safe(
                                    conflict.get("priority")
                                    or conflict.get("response_priority_level")
                                ),
                                _safe(
                                    conflict.get("event_name")
                                    or conflict.get("event_type")
                                    or conflict.get("financial_event_id")
                                ),
                                (
                                    f"{_safe(conflict.get('event_date'))}<br/>"
                                    f"gap {_safe(conflict.get('gap_days'))}일"
                                ),
                                _safe(conflict.get("reason") or conflict.get("impact_level")),
                            ]
                            for conflict in conflicts
                            if isinstance(conflict, dict)
                        ],
                    ],
                    font=font,
                    widths=[28 * mm, 42 * mm, 35 * mm, 56 * mm],
                )
            )
        data_actions = item.get("data_actions", [])
        if data_actions:
            story.append(
                Paragraph(
                    "<br/>".join(
                        (
                            f"• {_safe(action.get('message') or action.get('code'))}"
                            if isinstance(action, dict)
                            else f"• {_safe(action)}"
                        )
                        for action in data_actions
                    ),
                    body,
                )
            )
        if not conflicts and not data_actions:
            story.append(Paragraph("추가 충돌 또는 데이터 확인 신호가 없습니다.", body))

    story.extend(
        [
            Paragraph("근거·감사", heading),
            Paragraph(
                (
                    f"payload_hash={_safe(payload['payload_hash'])}<br/>"
                    "본 문서는 저장된 MonitoringRun과 Specialist Tool 결과를 재계산 없이 "
                    "요약한 내부 운영 보고서입니다."
                ),
                small,
            ),
            Spacer(1, 7 * mm),
        ]
    )
    document.build(
        story,
        onFirstPage=_daily_monitoring_footer,
        onLaterPages=_daily_monitoring_footer,
    )


def _customer_briefing_story(
    payload: dict[str, Any],
    styles: dict[str, ParagraphStyle],
) -> list[Any]:
    case = payload["case"]
    shipment = payload.get("shipment") or {}
    terms = payload.get("payment_terms") or {}
    numbers = _risk_numbers(payload)
    risk = numbers["risk"]
    conflict = numbers["conflict"]
    b_l_missing = not shipment.get("bl_no")
    anchor_missing = not shipment.get("on_board_date")

    if numbers["exceeded_days"]:
        risk_title = (
            f"안전한 마지막 본선적재일을 {numbers['exceeded_days']}일 초과했습니다"
        )
        risk_body = (
            f"B/L과 실제 본선적재일이 아직 확인되지 않았습니다. 현재 금융 일정 기준으로는 "
            f"{_display_date(numbers['latest_safe_anchor'])}까지 본선적재가 확인되어야 "
            f"{_display_date(numbers['event_date'])}의 금융 일정을 안전하게 준비할 수 있었습니다. "
            "이는 선제 위험 경고입니다. 아직 확정 연체가 아닙니다."
        )
    elif numbers["revised_due"]:
        risk_title = "선적 지연으로 금융 일정 충돌 가능성이 확인되었습니다"
        risk_body = (
            f"변경 예상 수출대금일은 {_display_date(numbers['revised_due'])}이며, 연결된 금융 일정은 "
            f"{_display_date(numbers['event_date'])}입니다. 실제 선적일과 B/L 발급일을 확인한 뒤 "
            "대체 자금과 만기 조정 가능성을 함께 검토해야 합니다."
        )
    else:
        risk_title = "금융 일정 확인이 필요합니다"
        risk_body = (
            "지급기준일 또는 실제 선적일이 아직 확정되지 않아 최종 지급일을 확정할 수 없습니다. "
            "선사와 실제 선적 사실 및 B/L 발급 상태를 먼저 확인해 주세요."
        )

    doc_rows = _doc_status_rows(payload)
    story: list[Any] = [
        Paragraph("고객용 거래 위험 안내서", styles["title"]),
        Paragraph(
            (
                f"{_safe(case.get('company'))} · 거래 {_safe(case.get('case_id'))} · "
                f"기준일 {_display_date(risk.get('as_of_date'))}"
            ),
            styles["subtitle"],
        ),
        Spacer(1, 5 * mm),
        _briefing_callout(
            risk_title,
            risk_body,
            styles=styles,
            background=KB_PALE_RED if conflict else KB_PALE_YELLOW,
            accent=KB_RED if conflict else KB_ORANGE,
        ),
        Paragraph("현재 거래 요약", styles["heading"]),
        _briefing_table(
            [
                ["구분", "확인 내용"],
                ["거래 상대방", _safe(case.get("counterparty"))],
                ["Invoice", _safe(case.get("invoice_no"))],
                ["Booking", _safe(shipment.get("booking_no"))],
                ["수출 품목", _safe(case.get("goods"))],
                [
                    "결제조건",
                    (
                        f"본선적재일 후 {_safe(terms.get('tenor_days'))}일 "
                        f"({_safe(terms.get('raw_text'))})"
                    ),
                ],
            ],
            styles=styles,
            widths=[42 * mm, 134 * mm],
        ),
        Paragraph("문서 보유 현황", styles["heading"]),
        _briefing_table(
            [["문서", "현재 상태"], *doc_rows],
            styles=styles,
            widths=[105 * mm, 71 * mm],
        ),
    ]

    if b_l_missing or anchor_missing:
        story.extend(
            [
                Spacer(1, 2 * mm),
                Paragraph(
                    (
                        "Booking과 Commercial Invoice는 확인되었지만 B/L과 본선적재일은 아직 "
                        "확인되지 않았습니다. B/L 부재만으로 미선적이나 연체를 확정하지는 않습니다."
                    ),
                    styles["body"],
                ),
            ]
        )

    metric_items = [
        [str(shipment.get("etd") or "미확정"), "예정 출항일"],
        [str(numbers["planned_due"] or "미확정"), "계획 지급일"],
        [str(numbers["event_date"] or "미확정"), "금융 일정일"],
        [
            str(numbers["latest_safe_anchor"] or numbers["revised_due"] or "미확정"),
            "최종 안전 본선적재일" if numbers["latest_safe_anchor"] else "변경 예상 지급일",
        ],
    ]
    story.extend(
        [
            Paragraph("핵심 일정", styles["heading"]),
            _metric_cards(metric_items, styles=styles, color=KB_YELLOW),
            Spacer(1, 3 * mm),
            Paragraph(
                (
                    f"예정 출항일은 <b>{_display_date(shipment.get('etd'))}</b>이었습니다. "
                    f"결제조건은 <b>{_humanize_anchor(terms.get('anchor_type_effective'))} 후 "
                    f"{_safe(terms.get('tenor_days'))}일</b>이며, 연결된 "
                    f"{_safe(conflict.get('event_name') or conflict.get('event_type') or '금융 일정')} 일정은 "
                    f"<b>{_display_date(numbers['event_date'])}</b>입니다."
                ),
                styles["body_large"],
            ),
            Paragraph("권장 조치", styles["heading"]),
            _briefing_callout(
                "지금 확인할 사항",
                (
                    "1. 선사에 실제 선적일과 B/L 발급 예정일을 확인합니다.<br/>"
                    "2. 대출 만기 조정 또는 단기 운전자금 확보 가능성을 검토합니다.<br/>"
                    "3. 수출채권 매입·팩토링 등 회수 지연 대응 수단을 검토합니다.<br/>"
                    "4. 실제 선적일이 확인되면 금융 일정을 즉시 다시 계산합니다."
                ),
                styles=styles,
                background=KB_PALE_YELLOW,
                accent=KB_YELLOW,
            ),
            Paragraph("관련 KB 상품 후보", styles["heading"]),
        ]
    )
    story.extend(_product_story(payload, styles=styles, rm=False))
    story.extend(
        [
            Paragraph("KB 무역금융 상담 안내", styles["heading"]),
            _briefing_callout(
                "KB 무역금융 상담사와 다음 대응을 준비해 보세요",
                (
                    "본 안내서와 실제 선적일·B/L 발급일을 함께 준비하면 만기 조정, 운전자금, "
                    "수출채권 활용 가능성을 더 빠르게 상담할 수 있습니다. 상품 이용 가능 여부와 "
                    "한도·금리·승인은 KB 상담 및 심사 후 확정됩니다."
                ),
                styles=styles,
                background=KB_PALE_YELLOW,
                accent=KB_DEEP_YELLOW,
            ),
            Spacer(1, 4 * mm),
            Paragraph(
                (
                    "본 문서는 저장된 거래·선적·금융 일정에 따른 선제 안내입니다. B/L 미수령과 "
                    "안전 기준 초과만으로 연체, 부도 또는 상품 이용 가능성을 확정하지 않습니다."
                ),
                styles["small"],
            ),
        ]
    )
    return story


def _rm_briefing_story(
    payload: dict[str, Any],
    styles: dict[str, ParagraphStyle],
) -> list[Any]:
    case = payload["case"]
    shipment = payload.get("shipment") or {}
    terms = payload.get("payment_terms") or {}
    numbers = _risk_numbers(payload)
    risk = numbers["risk"]
    conflict = numbers["conflict"]
    inventory = _document_inventory(payload)
    links = [item for item in payload.get("financial_links", []) if isinstance(item, dict)]
    link = next(
        (
            item
            for item in links
            if item.get("transaction_event_link_id")
            == conflict.get("transaction_event_link_id")
        ),
        links[0] if links else {},
    )
    dependency_scope = str(link.get("dependency_scope") or "UNKNOWN").upper()
    dependency_display = (
        dependency_scope if dependency_scope in {"FULL", "PARTIAL"} else "확인 필요"
    )
    linked_amount_display = (
        _display_amount(link.get("linked_amount"), link.get("linked_currency"))
        if link.get("linked_amount") is not None and link.get("linked_currency")
        else "확인 필요"
    )
    alternative_funds = str(link.get("alternative_funds_status") or "UNKNOWN").upper()
    alternative_display = (
        alternative_funds if alternative_funds not in {"", "UNKNOWN"} else "확인 필요"
    )

    headline = (
        f"B/L 미수령 · 안전 Anchor {numbers['exceeded_days']}일 초과 · {numbers['priority']}"
        if numbers["exceeded_days"]
        else f"선적 일정 변경 위험 · 우선순위 {numbers['priority']}"
    )
    story: list[Any] = [
        Paragraph("KB 직원용 거래 위험 검토서", styles["title"]),
        Paragraph(
            (
                f"{_safe(case.get('company'))} · Case {_safe(case.get('case_id'))} · "
                f"기준일 {_display_date(risk.get('as_of_date'))}"
            ),
            styles["subtitle"],
        ),
        Spacer(1, 5 * mm),
        _briefing_callout(
            headline,
            (
                "고객 안내 전 실제 선적일·B/L 발급 상태와 거래-금융이벤트 의존 범위를 확인해야 "
                "합니다. 현재 결과는 선제 충돌 신호이며 확정 연체 판정이 아닙니다."
            ),
            styles=styles,
            background=KB_PALE_RED,
            accent=KB_RED,
        ),
        Paragraph("업무 식별정보", styles["heading"]),
        _briefing_table(
            [
                ["항목", "값"],
                ["Company ID", _safe(case.get("company_id") or risk.get("company_id"))],
                ["Case ID", _safe(case.get("case_id"))],
                [
                    "Transaction ID",
                    _safe(case.get("transaction_id") or risk.get("transaction_id")),
                ],
                ["Calculation ID", _safe(risk.get("calculation_id"))],
                ["Basis Version", _safe(payload.get("basis_version"))],
            ],
            styles=styles,
            widths=[48 * mm, 128 * mm],
        ),
        Paragraph("문서 및 선적 증거", styles["heading"]),
        _briefing_table(
            [
                ["문서", "상태", "근거 ID"],
                [
                    "Booking Confirmation",
                    "보유",
                    _safe(inventory.get("BOOKING_CONFIRMATION", {}).get("document_id")),
                ],
                [
                    "Commercial Invoice",
                    "보유",
                    _safe(inventory.get("COMMERCIAL_INVOICE", {}).get("document_id")),
                ],
                [
                    "B/L (Bill of Lading)",
                    "미수령" if not shipment.get("bl_no") else "보유",
                    _safe(inventory.get("BILL_OF_LADING", {}).get("document_id")),
                ],
                ["B/L No.", _safe(shipment.get("bl_no")), "Shipment projection"],
                ["On Board Date", _safe(shipment.get("on_board_date")), "Shipment projection"],
                ["Booking ETD", _safe(shipment.get("etd")), "Booking fact"],
            ],
            styles=styles,
            widths=[58 * mm, 35 * mm, 83 * mm],
        ),
        Paragraph("지급조건 파싱 결과", styles["heading"]),
        _briefing_table(
            [
                ["항목", "검증 결과"],
                ["원문", _safe(terms.get("raw_text"))],
                ["Anchor", _safe(terms.get("anchor_type_effective"))],
                ["Tenor", f"{_safe(terms.get('tenor_days'))}일"],
                ["Day Type", _safe(terms.get("day_type_effective"))],
                ["Calculation Gate", "허용" if terms.get("verified") else "추가 검증 필요"],
            ],
            styles=styles,
            widths=[48 * mm, 128 * mm],
        ),
        Paragraph("위험 계산 근거", styles["heading"]),
        _metric_cards(
            [
                [str(numbers["planned_due"] or "미확정"), "계획 지급일"],
                [str(numbers["event_date"] or "미확정"), "금융 일정일"],
                [
                    str(numbers["latest_safe_anchor"] or numbers["revised_due"] or "미확정"),
                    "최종 안전 Anchor" if numbers["latest_safe_anchor"] else "변경 지급일",
                ],
                [
                    f"{numbers['exceeded_days']}일" if numbers["exceeded_days"] else str(numbers["priority"]),
                    "기준 초과" if numbers["exceeded_days"] else "우선순위",
                ],
            ],
            styles=styles,
            color=KB_YELLOW,
        ),
        Spacer(1, 3 * mm),
        _briefing_table(
            [
                ["항목", "값"],
                ["충돌 이벤트", _safe(conflict.get("event_name") or conflict.get("event_type"))],
                [
                    "금융기관",
                    _safe(
                        link.get("financial_institution")
                        or conflict.get("financial_institution")
                        or "확인 필요"
                    ),
                ],
                ["금액", _display_amount(conflict.get("amount"), conflict.get("currency"))],
                ["충돌 원인", _safe(conflict.get("conflict_origin"))],
                ["우선순위", _safe(numbers["priority"])],
                ["판정 근거", _safe(conflict.get("reason"))],
            ],
            styles=styles,
            widths=[48 * mm, 128 * mm],
        ),
        Paragraph("거래-금융일정 연결 확인", styles["heading"]),
        _briefing_table(
            [
                ["항목", "현재 저장값 / 조치"],
                ["Link ID", _safe(link.get("transaction_event_link_id"))],
                ["Link Type", _safe(link.get("link_type"))],
                ["의존 범위", _safe(dependency_display)],
                ["연결 금액·통화", _safe(linked_amount_display)],
                ["대체자금 상태", _safe(alternative_display)],
                ["적용 환율", "DB 미보유 · 상담 시 확인 필요"],
            ],
            styles=styles,
            widths=[48 * mm, 128 * mm],
        ),
        Paragraph("상담 준비 체크리스트", styles["heading"]),
        _briefing_callout(
            "고객에게 확인할 운영 정보",
            (
                f"□ 의존 범위: {_safe(dependency_display)} — 전액(FULL)/일부(PARTIAL) 확인<br/>"
                f"□ 연결 금액·통화: {_safe(linked_amount_display)} — 적용 환율 함께 확인<br/>"
                f"□ 대체 운전자금 또는 상환 재원: {_safe(alternative_display)}<br/>"
                "□ 선사에 실제 선적일·본선적재일·B/L 발급 예정일 확인<br/>"
                "□ 실제 Anchor 확인 후 위험 계산 재실행"
            ),
            styles=styles,
            background=KB_PALE_YELLOW,
            accent=KB_YELLOW,
        ),
        Paragraph("관련 KB 상품 후보와 원문 근거", styles["heading"]),
    ]
    story.extend(_product_story(payload, styles=styles, rm=True))
    story.extend(
        [
            Paragraph("고객 상담 질문 목록", styles["heading"]),
            _briefing_table(
                [
                    ["순서", "질문"],
                    ["1", "선사로부터 실제 선적일과 B/L 발급 예정일을 확인하셨습니까?"],
                    ["2", "이 금융 이벤트는 해당 수출대금에 전액 의존합니까, 일부만 의존합니까?"],
                    ["3", "연결 금액과 통화, 적용 환율은 어떻게 됩니까?"],
                    ["4", "만기 전 활용 가능한 대체 자금 또는 상환 재원이 있습니까?"],
                    ["5", "대출 만기 조정·운전자금·수출채권 활용 중 우선 검토 목적은 무엇입니까?"],
                ],
                styles=styles,
                widths=[18 * mm, 158 * mm],
            ),
            Paragraph("감사 및 산출물 기준", styles["heading"]),
            Paragraph(
                (
                    f"source_kind={_safe(risk.get('source_kind'))}<br/>"
                    f"source_ref_id={_safe(risk.get('source_ref_id'))}<br/>"
                    f"priority_rule_version={_safe(risk.get('priority_rule_version'))}<br/>"
                    f"input_fingerprint={_safe(risk.get('input_fingerprint'))}<br/>"
                    f"payload_hash={_safe(payload.get('payload_hash'))}"
                ),
                styles["small"],
            ),
        ]
    )
    return story


def _render_briefing_pdf(
    path: Path,
    *,
    audience: str,
    payload: dict[str, Any],
) -> None:
    """Render the customer-friendly or RM-review briefing from one frozen payload."""
    font = _register_font()
    styles = _briefing_styles(font)
    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=17 * mm,
        leftMargin=17 * mm,
        topMargin=28 * mm,
        bottomMargin=20 * mm,
        title=(
            "KB TradeFlow Twin 고객용 거래 위험 안내서"
            if audience == "CUSTOMER"
            else "KB TradeFlow Twin 직원용 거래 위험 검토서"
        ),
        author="KB TradeFlow Twin",
    )
    story = (
        _customer_briefing_story(payload, styles)
        if audience == "CUSTOMER"
        else _rm_briefing_story(payload, styles)
    )
    page_callback = lambda canvas, doc: _kb_briefing_page(canvas, doc, audience)
    document.build(
        story,
        onFirstPage=page_callback,
        onLaterPages=page_callback,
    )


class ReportService:
    """Persist internal monitoring reports and consent-gated case briefings."""

    def __init__(self, session: Session) -> None:
        self.session = session

    def build_daily_monitoring_payload(
        self,
        monitoring_run_id: str,
        monitoring_payload: dict[str, Any],
    ) -> dict[str, Any]:
        """Freeze one reviewed monitoring result without recalculating its risk."""
        run = self.session.get(MonitoringRun, monitoring_run_id)
        if run is None:
            raise KeyError(monitoring_run_id)
        supplied_run_id = monitoring_payload.get("monitoring_run_id")
        if supplied_run_id and str(supplied_run_id) != monitoring_run_id:
            raise ValueError("monitoring_payload belongs to another MonitoringRun")
        supplied_date = monitoring_payload.get("as_of_date")
        if supplied_date and str(supplied_date) != run.as_of_date.isoformat():
            raise ValueError("monitoring_payload as_of_date does not match MonitoringRun")

        nested = monitoring_payload.get("final_response")
        source = dict(monitoring_payload)
        if isinstance(nested, dict):
            source.update(nested)
        ranked_cases = _daily_case_details(source)
        priority_summary = source.get("priority_summary", {})
        if not isinstance(priority_summary, dict):
            priority_summary = {}
        if not priority_summary:
            priority_counts: dict[str, int] = {}
            for item in ranked_cases:
                label = str(item.get("highest_priority") or "UNRANKED")
                priority_counts[label] = priority_counts.get(label, 0) + 1
            priority_summary = {
                "total_ranked": len(ranked_cases),
                "cases_with_conflicts": sum(
                    1 for item in ranked_cases if item["conflict_count"] > 0
                ),
                "total_conflicts": sum(item["conflict_count"] for item in ranked_cases),
                "by_highest_priority": {
                    label: priority_counts[label]
                    for label in sorted(priority_counts, key=_priority_sort_key)
                },
                "ordered_case_ids": [item["case_id"] for item in ranked_cases],
            }

        priorities = [
            str(item["highest_priority"]) for item in ranked_cases if item.get("highest_priority")
        ]
        highest_priority = min(priorities, key=_priority_sort_key) if priorities else None
        risk_case_count = sum(
            1
            for item in ranked_cases
            if item["conflict_count"] > 0 or bool(item.get("data_actions"))
        )
        trace = run.trace_json if isinstance(run.trace_json, dict) else {}
        specialist_call_count = int(
            source.get("specialist_call_count")
            or trace.get("specialist_call_count")
            or (run.selected_count * 2)
        )
        if risk_case_count:
            priority_text = highest_priority or "미분류"
            summary_text = (
                f"{run.as_of_date.isoformat()} 기준 후보 {run.candidate_count}건을 점검했고 "
                f"{risk_case_count}건에서 금융 일정 충돌 또는 데이터 확인 필요 신호를 "
                f"확인했습니다. 최고 우선순위는 {priority_text}입니다."
            )
        else:
            summary_text = (
                f"{run.as_of_date.isoformat()} 기준 후보 {run.candidate_count}건을 점검했으며 "
                "표시할 금융 일정 충돌 또는 데이터 확인 필요 신호가 없습니다."
            )
        summary = {
            "monitoring_run_id": run.monitoring_run_id,
            "as_of_date": run.as_of_date.isoformat(),
            "candidate_count": run.candidate_count,
            "selected_count": run.selected_count,
            "result_count": run.result_count,
            "specialist_call_count": specialist_call_count,
            "new_alert_count": run.new_alert_count,
            "risk_case_count": risk_case_count,
            "highest_priority": highest_priority,
            "priority_summary": priority_summary,
            "ordered_case_ids": [item["case_id"] for item in ranked_cases],
            "summary_text": summary_text,
        }
        payload: dict[str, Any] = {
            "monitoring_run_id": run.monitoring_run_id,
            "as_of_date": run.as_of_date.isoformat(),
            "summary": summary,
            "ranked_cases": ranked_cases,
        }
        payload["payload_hash"] = _payload_hash(payload)
        return payload

    def generate_daily_monitoring_report(
        self,
        monitoring_run_id: str,
        monitoring_payload: dict[str, Any],
    ) -> dict[str, Any]:
        """Idempotently render and persist one internal daily monitoring PDF."""
        payload = self.build_daily_monitoring_payload(
            monitoring_run_id,
            monitoring_payload,
        )
        payload_hash = str(payload["payload_hash"])
        dedup_key = _payload_hash(
            {
                "monitoring_run_id": monitoring_run_id,
                "payload_hash": payload_hash,
            }
        )
        daily_report_id = _stable_id("DMR", dedup_key)
        asset = REPORT_DIR / f"daily_monitoring_{payload['as_of_date']}_{payload_hash[:12]}.pdf"
        existing = self.session.get(DailyMonitoringReport, daily_report_id)
        if existing is not None and asset.is_file():
            current_hash = _hash_file(asset)
            if existing.content_hash == current_hash:
                summary = dict(existing.summary_json)
                return {
                    "status": "READY",
                    "daily_report_id": existing.daily_report_id,
                    "monitoring_run_id": existing.monitoring_run_id,
                    "as_of_date": existing.as_of_date.isoformat(),
                    "summary": summary,
                    "daily_summary": summary.get("summary_text"),
                    "highest_priority": existing.highest_priority,
                    "risk_case_count": existing.risk_case_count,
                    "asset_path": existing.asset_path,
                    "pdf_path": existing.asset_path,
                    "content_hash": existing.content_hash,
                    "payload_hash": existing.payload_hash,
                    "deduplicated": True,
                }

        REPORT_DIR.mkdir(parents=True, exist_ok=True)
        _render_daily_monitoring_pdf(asset, payload)
        content_hash = _hash_file(asset)
        summary = dict(payload["summary"])
        if existing is None:
            report = DailyMonitoringReport(
                daily_report_id=daily_report_id,
                monitoring_run_id=monitoring_run_id,
                as_of_date=date.fromisoformat(str(payload["as_of_date"])),
                summary_json=summary,
                highest_priority=summary.get("highest_priority"),
                risk_case_count=int(summary["risk_case_count"]),
                asset_path=str(asset),
                content_hash=content_hash,
                payload_hash=payload_hash,
                dedup_key=dedup_key,
            )
            self.session.add(report)
        else:
            report = existing
            report.summary_json = summary
            report.highest_priority = summary.get("highest_priority")
            report.risk_case_count = int(summary["risk_case_count"])
            report.asset_path = str(asset)
            report.content_hash = content_hash
            report.payload_hash = payload_hash
        self.session.flush()
        return {
            "status": "READY",
            "daily_report_id": report.daily_report_id,
            "monitoring_run_id": report.monitoring_run_id,
            "as_of_date": report.as_of_date.isoformat(),
            "summary": summary,
            "daily_summary": summary["summary_text"],
            "highest_priority": report.highest_priority,
            "risk_case_count": report.risk_case_count,
            "asset_path": report.asset_path,
            "pdf_path": report.asset_path,
            "content_hash": report.content_hash,
            "payload_hash": report.payload_hash,
            "deduplicated": existing is not None,
        }

    def build_briefing_payload(
        self,
        case_id: str,
        *,
        product_options: list[dict[str, Any]] | None = None,
        as_of_date: date | None = None,
        rm_inbox: str | None = None,
    ) -> dict[str, Any]:
        """Freeze workbook, case facts, product evidence, and one risk snapshot."""
        trade_case = self.session.get(TradeCase, case_id)
        if trade_case is None:
            raise KeyError(case_id)
        shipment = self.session.scalar(select(Shipment).where(Shipment.case_id == case_id))
        obligation = self.session.scalar(
            select(PaymentObligation)
            .where(PaymentObligation.case_id == case_id)
            .order_by(PaymentObligation.tranche_index)
        )
        financial_links = list(
            self.session.scalars(
                select(TransactionFinancialEventLink)
                .where(TransactionFinancialEventLink.case_id == case_id)
                .order_by(TransactionFinancialEventLink.transaction_event_link_id)
            )
        )
        linked_events = {
            link.financial_event_id: self.session.get(FinancialEvent, link.financial_event_id)
            for link in financial_links
        }
        financial_reference = read_transaction_financial_reference(self.session, trade_case)
        try:
            risk_snapshot = FinancialExposureService(self.session).get_risk_snapshot(
                case_id, as_of_date
            )
        except (KeyError, ValueError):
            risk_snapshot = {
                "case_id": case_id,
                "company_id": trade_case.company_id,
                "transaction_id": trade_case.transaction_id,
                "calculation_id": None,
                "basis_version": trade_case.basis_version,
                "as_of_date": as_of_date.isoformat() if as_of_date else None,
                "source_kind": "NO_PERSISTED_RISK_SNAPSHOT",
                "source_ref_id": None,
                "expected_receipt": {
                    "date": None,
                    "amount": financial_reference["amount"],
                    "currency": financial_reference["currency"],
                },
                "highest_priority": None,
                "conflict_count": 0,
                "conflicts": [],
                "data_actions": [
                    {
                        "code": "RISK_SNAPSHOT_REQUIRED",
                        "message": (
                            "지급기준일과 금융 일정 계산 결과가 없어 "
                            "충돌 표시 전에 risk scan이 필요합니다."
                        ),
                    }
                ],
                "priority_rule_version": None,
                "input_fingerprint": None,
            }
        blueprint = load_briefing_blueprint()
        documents = list(self.session.scalars(select(Document).where(Document.case_id == case_id)))
        evidence = [
            {
                "label": f"문서 {document.doc_type}",
                "reference": (
                    f"{document.document_id} · sha256:{document.sha256} · "
                    f"object:{Path(document.object_path).name}"
                ),
            }
            for document in documents
        ]
        if risk_snapshot.get("calculation_id"):
            evidence.append(
                {
                    "label": "공유 risk snapshot",
                    "reference": (
                        f"{risk_snapshot['calculation_id']} · "
                        f"fingerprint:{risk_snapshot.get('input_fingerprint')}"
                    ),
                }
            )
        for option in product_options or []:
            for citation in option.get("citations", []):
                evidence.append(
                    {
                        "label": f"상품 {option.get('canonical_name')}",
                        "reference": (
                            f"{citation.get('source_file')} p.{citation.get('page')} · "
                            f"sha256:{citation.get('source_sha256')}"
                        ),
                    }
                )
        basis = str(risk_snapshot.get("basis_version") or trade_case.basis_version)
        payload: dict[str, Any] = {
            "case": {
                "case_id": trade_case.case_id,
                "company_id": trade_case.company_id,
                "transaction_id": trade_case.transaction_id,
                "company": trade_case.company,
                "counterparty": trade_case.counterparty,
                "invoice_no": trade_case.invoice_no,
                "contract_amount": (
                    risk_snapshot.get("expected_receipt", {}).get("amount")
                    if (
                        isinstance(risk_snapshot.get("expected_receipt"), dict)
                        and risk_snapshot.get("expected_receipt", {}).get("amount") is not None
                    )
                    else financial_reference["amount"]
                ),
                "currency": trade_case.currency,
                "goods": trade_case.goods,
                "status": trade_case.status,
            },
            "shipment": (
                {
                    "booking_no": shipment.booking_no,
                    "bl_no": shipment.bl_no,
                    "etd": shipment.etd.isoformat() if shipment.etd else None,
                    "on_board_date": (
                        shipment.on_board_date.isoformat() if shipment.on_board_date else None
                    ),
                    "status": shipment.status,
                    "departure_status": (
                        "CONFIRMED_DEPARTED" if shipment.on_board_date is not None else "UNKNOWN"
                    ),
                }
                if shipment
                else None
            ),
            "payment_terms": (
                {
                    "raw_text": obligation.raw_text,
                    "anchor_type_effective": obligation.anchor_type_effective,
                    "tenor_days": obligation.tenor_days,
                    "day_type_effective": obligation.day_type_effective,
                    "verified": obligation.verified,
                }
                if obligation
                else {}
            ),
            "financial_links": [
                {
                    "transaction_event_link_id": link.transaction_event_link_id,
                    "financial_event_id": link.financial_event_id,
                    "event_name": (
                        linked_events[link.financial_event_id].event_name
                        if linked_events.get(link.financial_event_id)
                        else None
                    ),
                    "financial_institution": (
                        linked_events[link.financial_event_id].institution
                        if linked_events.get(link.financial_event_id)
                        else None
                    ),
                    "link_type": link.link_type,
                    "link_status": link.link_status,
                    "dependency_scope": link.dependency_scope,
                    "linked_amount": link.linked_amount,
                    "linked_currency": link.linked_currency,
                    "alternative_funds_status": link.alternative_funds_status,
                }
                for link in financial_links
            ],
            "risk_snapshot": risk_snapshot,
            "documents": [
                {
                    "document_id": document.document_id,
                    "doc_type": document.doc_type,
                    "sha256": document.sha256,
                    "object_name": Path(document.object_path).name,
                }
                for document in documents
            ],
            "product_options": product_options or [],
            "evidence": evidence,
            "blueprint": blueprint,
            "basis_version": basis,
            "rm_inbox": rm_inbox,
        }
        payload["payload_hash"] = _payload_hash(payload)
        return payload

    def render_report(
        self,
        *,
        briefing_payload: dict[str, Any],
        audience: str,
        consent: bool,
    ) -> dict[str, Any]:
        """Render and persist one audience from the frozen briefing payload."""
        if not consent:
            raise ValueError("Explicit report consent is required")
        normalized_audience = audience.strip().upper()
        if normalized_audience not in {"CUSTOMER", "RM"}:
            raise ValueError("audience must be CUSTOMER or RM")
        case_id = str(briefing_payload["case"]["case_id"])
        if self.session.get(TradeCase, case_id) is None:
            raise KeyError(case_id)
        safe_case_id = re.sub(r"[^A-Za-z0-9._-]+", "-", case_id).strip("-")
        if not safe_case_id:
            raise ValueError("case_id cannot be represented as a report filename")
        basis = str(briefing_payload["basis_version"])
        REPORT_DIR.mkdir(parents=True, exist_ok=True)
        suffix = "customer" if normalized_audience == "CUSTOMER" else "rm"
        asset = REPORT_DIR / f"{safe_case_id}_{suffix}.pdf"
        _render_briefing_pdf(
            asset,
            audience=normalized_audience,
            payload=briefing_payload,
        )
        report_id = _stable_id(
            "RPT",
            (f"{case_id}:{basis}:{normalized_audience}:{briefing_payload['payload_hash']}"),
        )
        report = self.session.get(Report, report_id)
        if report is None:
            report = Report(
                report_id=report_id,
                case_id=case_id,
                audience=normalized_audience,
                basis_version=basis,
                asset_path=str(asset),
                content_hash=_hash_file(asset),
            )
            self.session.add(report)
        else:
            report.asset_path = str(asset)
            report.content_hash = _hash_file(asset)
        self.session.flush()
        return {
            "report_id": report_id,
            "case_id": case_id,
            "audience": normalized_audience,
            "basis_version": basis,
            "asset_path": str(asset),
            "content_hash": report.content_hash,
            "payload_hash": briefing_payload["payload_hash"],
            "briefing_source_sha256": briefing_payload["blueprint"]["source_sha256"],
        }
