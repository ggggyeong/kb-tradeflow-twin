"""Read the original two-sheet calendar; optionally verify document numbers on sheet 3."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from app.schemas.portfolio import PortfolioFinancialEvent

EVENT_SHEET = "1.금융이벤트"
LINK_SHEET = "2.거래연결"
TRADE_SHEET = "3.거래정보"
SCENARIO_NAMES = {
    "WORKING_CAPITAL_LOAN_MATURITY": "운전자금 대출 만기",
    "FX_FORWARD_MATURITY": "선물환 결제일",
    "SUPPLIER_PAYMENT": "공급자 지급일",
}
LINK_TYPES = {
    "WORKING_CAPITAL_LOAN_MATURITY": {"LOAN_REPAYMENT_SOURCE"},
    "FX_FORWARD_MATURITY": {"FX_SETTLEMENT_SOURCE"},
    # Retain the original demo_2 workbook's name without broadening other scenarios.
    "SUPPLIER_PAYMENT": {"SUPPLIER_PAYMENT_SOURCE", "EXPECTED_EXPORT_RECEIPT"},
}


@dataclass(frozen=True)
class FinancialCalendar:
    events: list[PortfolioFinancialEvent]
    document_references: dict[str, str]
    source_file: str
    warnings: list[str]


def _rows(workbook: Any, sheet: str, required: set[str]) -> list[dict[str, Any]]:
    if sheet not in workbook.sheetnames:
        raise ValueError(f"필수 엑셀 시트 누락: {sheet}")
    ws = workbook[sheet]
    # Read-only worksheets may omit or understate dimensions; inspect actual cells.
    # The ZIP's uncompressed-size limit bounds this scan before loading the workbook.
    if hasattr(ws, "reset_dimensions"):
        ws.reset_dimensions()
        ws.calculate_dimension(force=True)
    if (ws.max_row or 0) > 1001 or (ws.max_column or 0) > 20:
        raise ValueError("엑셀은 시트당 1,000개 데이터 행·20개 열 이내여야 합니다.")
    # Some producers omit worksheet dimensions. Bound the scan independently of metadata.
    scanned = list(ws.iter_rows(max_row=1002, max_col=21))
    if (len(scanned) > 1001 and any(c.value is not None for c in scanned[1001])) or any(
        row[-1].value is not None for row in scanned
    ):
        raise ValueError("엑셀 입력 행·열 제한을 초과했습니다.")
    values = [row[:20] for row in scanned[:1001]]
    if not values:
        raise ValueError(f"엑셀 시트가 비어 있습니다: {sheet}")
    headers = [str(c.value).strip() if c.value is not None else "" for c in values[0]]
    named = [h for h in headers if h]
    if len(named) != len(set(named)) or not required.issubset(named):
        raise ValueError(f"엑셀 필수 열 누락 또는 중복: {sheet}")
    result = []
    for row in values[1:]:
        if any(c.data_type == "f" for c in row):
            raise ValueError("입력 엑셀은 수식 대신 확정된 값을 사용해야 합니다.")
        item = {h: c.value for h, c in zip(headers, row, strict=True) if h}
        if any(v is not None for v in item.values()):
            result.append(item)
    return result


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _event_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None or value == "":
        return None
    # Bare Excel serials and locale-ambiguous strings are not guessed.
    return date.fromisoformat(_text(value))


def read_financial_calendar(path: Path, transaction_id: str) -> FinancialCalendar:
    from openpyxl import load_workbook

    if path.suffix.lower() != ".xlsx" or path.stat().st_size > 5_000_000:
        raise ValueError("금융일정은 5MB 이하 .xlsx 파일이어야 합니다.")
    with ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > 30_000_000:
            raise ValueError("압축 해제 크기가 너무 큰 엑셀입니다.")
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        events = _rows(workbook, EVENT_SHEET, {"event_id", "event_type_code", "event_date"})
        links = _rows(workbook, LINK_SHEET, {"transaction_id", "event_id", "link_status"})
        by_id = {}
        for event in events:
            event_id = _text(event["event_id"])
            if not event_id or event_id in by_id:
                raise ValueError("금융이벤트 ID가 비어 있거나 중복됩니다.")
            by_id[event_id] = event
        selected = [row for row in links if _text(row["transaction_id"]) == transaction_id]
        if not selected:
            raise ValueError("선택한 거래에 연결된 금융일정이 없습니다.")
        parsed = []
        seen = set()
        for link in selected:
            event_id = _text(link["event_id"])
            if event_id in seen or event_id not in by_id:
                raise ValueError("거래 연결의 이벤트 ID가 중복되거나 존재하지 않습니다.")
            seen.add(event_id)
            row = by_id[event_id]
            code = _text(row["event_type_code"])
            link_type = _text(link.get("link_type"))
            if link_type and link_type not in LINK_TYPES.get(code, set()):
                raise ValueError("거래 연결 유형과 금융이벤트 종류가 일치하지 않습니다.")
            parsed.append(
                PortfolioFinancialEvent.model_validate(
                    {
                        "event_id": event_id,
                        "scenario_code": code,
                        "event_name": _text(row.get("event_name"))
                        or SCENARIO_NAMES.get(code, code),
                        "event_date": _event_date(row["event_date"]),
                        "amount": row.get("amount"),
                        "currency": _text(row.get("currency")) or None,
                        "financial_institution": _text(row.get("financial_institution")) or None,
                        "link_status": _text(link["link_status"]) or "UNCONFIRMED",
                        "payment_purpose": _text(row.get("payment_purpose")) or "UNKNOWN",
                        "fx_direction": _text(row.get("fx_direction")) or "UNKNOWN",
                    }
                )
            )
        if len(parsed) > 3:
            raise ValueError("포트폴리오 분석은 한 거래의 금융일정 3개 이내를 지원합니다.")
        references: dict[str, str] = {}
        if TRADE_SHEET in workbook.sheetnames:
            trades = _rows(
                workbook, TRADE_SHEET, {"transaction_id", "invoice_no", "bl_no", "booking_no"}
            )
            matches = [r for r in trades if _text(r["transaction_id"]) == transaction_id]
            if len(matches) != 1:
                raise ValueError("거래정보에서 선택한 거래를 유일하게 확인할 수 없습니다.")
            references = {
                k: _text(matches[0][k])
                for k in ("invoice_no", "bl_no", "booking_no")
                if _text(matches[0][k])
            }
        warnings = (
            []
            if references
            else ["엑셀에 문서 번호 연결 정보가 없습니다. 대금 유입일 자동 계산은 보류합니다."]
        )
        return FinancialCalendar(parsed, references, path.name, warnings)
    finally:
        workbook.close()
