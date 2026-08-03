from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.domain_inputs.loaders.field_dictionary import sha256_file
from app.schemas.payment_terms import PaymentTermExample, PaymentTermsContract

EXPECTED_SHEETS = [
    "1.분해 컬럼 스펙",
    "2.결제조건 분류(핵심_부차)",
    "3.고객확인 질문",
    "4.예문 세트",
]
EXPECTED_FIRST_HEADERS = {
    "1.분해 컬럼 스펙": "column_name",
    "2.결제조건 분류(핵심_부차)": "그룹",
    "3.고객확인 질문": "column_name",
    "4.예문 세트": "그룹",
}


def _normalize_optional(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return None if text in {"", "-", "null", "None"} else text


def load_payment_terms_contract(path: Path) -> PaymentTermsContract:
    """Load the reviewed four-sheet PaymentTerms contract with no seed fallback."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    if workbook.sheetnames != EXPECTED_SHEETS:
        raise ValueError(f"PaymentTerms sheet mismatch: {workbook.sheetnames}")

    sheet_rows: dict[str, list[dict[str, Any]]] = {}
    for sheet_name in EXPECTED_SHEETS:
        sheet = workbook[sheet_name]
        all_rows = list(sheet.iter_rows(values_only=True))
        headers = [str(value or "") for value in all_rows[0]]
        if headers[0] != EXPECTED_FIRST_HEADERS[sheet_name]:
            raise ValueError(f"Invalid first header for {sheet_name}: {headers[0]!r}")
        data: list[dict[str, Any]] = []
        for values in all_rows[1:]:
            first = values[0]
            if first is None or str(first).startswith("["):
                continue
            data.append({headers[index]: value for index, value in enumerate(values)})
        sheet_rows[sheet_name] = data

    example_rows = sheet_rows["4.예문 세트"]
    if len(example_rows) != 21:
        raise ValueError(f"Expected 21 PaymentTerms examples, got {len(example_rows)}")
    examples = [
        PaymentTermExample(
            group=str(row["그룹"]),
            category=str(row["구분"]),
            raw_text=str(row["원문(raw_text)"]),
            payment_ratio=_normalize_optional(row["결제비율"]),
            payment_method=str(row["payment_method"]),
            credit_term_type=_normalize_optional(row["credit_term_type"]),
            lc_type=_normalize_optional(row["lc_type"]),
            anchor_type_extracted=str(row["anchor_type_extracted"]),
            anchor_type_effective=str(row["anchor_type_effective"]),
            tenor_days=int(row["tenor_days"]) if row["tenor_days"] is not None else None,
            day_type_extracted=str(row["day_type_extracted"]),
            day_type_effective=_normalize_optional(row["day_type_effective"]),
            advance_or_deferred=str(row["advance_or_deferred"]),
            calculation_allowed=str(row["calculation_allowed"]),
            document_stage=str(row["document_stage"]),
            real_world_status=str(row["real_world_status"]),
            reason=str(row["모호_특이 사유"] or ""),
            customer_question=_normalize_optional(row["고객 확인 질문"]),
        )
        for row in example_rows
    ]
    return PaymentTermsContract(
        source_sha256=sha256_file(path),
        source_path=str(path),
        sheet_rows=sheet_rows,
        examples=examples,
    )
