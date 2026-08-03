from __future__ import annotations

import hashlib
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.domain_inputs.loaders.field_dictionary import sha256_file
from app.schemas.financial_calendar import (
    FinancialCalendarContract,
    FinancialCalendarEventRow,
    FinancialCalendarLinkRow,
)

REQUIRED_SHEETS = ["1.금융이벤트", "2.거래연결"]
EVENT_HEADERS = [
    "event_id",
    "event_type_code",
    "event_date",
    "amount",
    "currency",
    "financial_institution",
]
LINK_HEADERS = [
    "transaction_id",
    "event_id",
    "link_type",
    "link_status",
]

EVENT_TYPE_NAMES = {
    "SUPPLIER_PAYMENT": "공급업체 지급",
    "WORKING_CAPITAL_LOAN_MATURITY": "운전자금대출 만기",
    "FX_FORWARD_MATURITY": "선물환 계약 만기",
}
EVENT_TYPE_CODES = set(EVENT_TYPE_NAMES)
LINK_STATUSES = {"CONFIRMED", "UNCONFIRMED", "NOT_LINKED"}


def _required_text(value: Any, field: str) -> str:
    text = str(value or "").strip()
    if not text:
        raise ValueError(f"{field} is required")
    return text


def _date(value: Any, field: str) -> date:
    if value is None or str(value).strip() == "":
        raise ValueError(f"{field} is required")
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError as exc:
        raise ValueError(f"{field} must be ISO-8601 YYYY-MM-DD: {value!r}") from exc


def _table(sheet: Worksheet, expected_headers: list[str]) -> list[dict[str, Any]]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Required sheet is empty: {sheet.title}")
    headers = [str(value or "").strip() for value in rows[0]]
    if headers != expected_headers:
        raise ValueError(
            f"Header mismatch for {sheet.title}: expected {expected_headers}, got {headers}"
        )
    return [
        dict(zip(headers, values, strict=True))
        for values in rows[1:]
        if any(value not in (None, "") for value in values)
    ]


def _is_kb_institution(value: str) -> bool:
    normalized = "".join(value.upper().split())
    return normalized.startswith("KB") or "KB국민은행" in normalized or "국민은행" in normalized


def _link_id(company_id: str, transaction_id: str, event_id: str) -> str:
    identity = f"{company_id}:{transaction_id}:{event_id}"
    return f"LNK-{hashlib.sha256(identity.encode()).hexdigest()[:20]}"


def _validate_contract(contract: FinancialCalendarContract) -> None:
    if not contract.company_id.strip():
        raise ValueError("company_id is required from the authenticated session")
    if not contract.events:
        raise ValueError("At least one financial event is required")
    if not contract.links:
        raise ValueError("At least one transaction link is required")

    event_ids = [row.event_id for row in contract.events]
    if len(set(event_ids)) != len(event_ids):
        raise ValueError("Duplicate event_id in financial calendar")

    link_pairs = [(row.transaction_id, row.event_id) for row in contract.links]
    if len(set(link_pairs)) != len(link_pairs):
        raise ValueError("Duplicate transaction_id/event_id link in financial calendar")

    known_events = set(event_ids)
    for event in contract.events:
        if event.event_type_code not in EVENT_TYPE_CODES:
            raise ValueError(f"Unsupported event_type_code: {event.event_type_code}")
        if event.amount <= 0:
            raise ValueError(f"Event amount must be positive: {event.event_id}")
        if len(event.currency) != 3 or not event.currency.isalpha():
            raise ValueError(f"currency must be a three-letter code: {event.event_id}")

    for link in contract.links:
        if link.event_id not in known_events:
            raise ValueError(
                f"Unknown event_id {link.event_id!r} for transaction {link.transaction_id!r}"
            )
        if link.link_status not in LINK_STATUSES:
            raise ValueError(f"Unsupported link_status: {link.link_status}")
        if not link.link_type:
            raise ValueError(f"link_type is required for transaction {link.transaction_id}")


def load_financial_calendar_contract(
    path: Path,
    *,
    company_id: str,
) -> FinancialCalendarContract:
    """Load the exact two-sheet, single-company financial-calendar contract.

    `company_id` is trusted execution context supplied by the authenticated
    session. It is intentionally absent from workbook rows and cannot be
    changed by the uploaded file.
    """
    session_company_id = _required_text(company_id, "company_id")
    workbook = load_workbook(path, read_only=True, data_only=True)
    if workbook.sheetnames != REQUIRED_SHEETS:
        raise ValueError(
            f"Financial calendar requires exactly {REQUIRED_SHEETS}, got {workbook.sheetnames}"
        )

    events: list[FinancialCalendarEventRow] = []
    for row in _table(workbook[REQUIRED_SHEETS[0]], EVENT_HEADERS):
        event_id = _required_text(row["event_id"], "event_id")
        event_type_code = _required_text(row["event_type_code"], "event_type_code").upper()
        institution = _required_text(row["financial_institution"], "financial_institution")
        try:
            amount = float(row["amount"])
        except (TypeError, ValueError) as exc:
            raise ValueError(f"amount must be numeric for event {event_id}") from exc
        events.append(
            FinancialCalendarEventRow(
                event_id=event_id,
                event_type_code=event_type_code,
                event_name=EVENT_TYPE_NAMES.get(event_type_code, event_type_code),
                event_date=_date(row["event_date"], f"event_date:{event_id}"),
                amount=amount,
                currency=_required_text(row["currency"], "currency").upper(),
                financial_institution=institution,
                is_kb_contract=_is_kb_institution(institution),
            )
        )

    links: list[FinancialCalendarLinkRow] = []
    for row in _table(workbook[REQUIRED_SHEETS[1]], LINK_HEADERS):
        transaction_id = _required_text(row["transaction_id"], "transaction_id")
        event_id = _required_text(row["event_id"], "event_id")
        links.append(
            FinancialCalendarLinkRow(
                link_id=_link_id(session_company_id, transaction_id, event_id),
                transaction_id=transaction_id,
                event_id=event_id,
                link_type=_required_text(row["link_type"], "link_type").upper(),
                link_status=_required_text(row["link_status"], "link_status").upper(),
                dependency_scope="UNKNOWN",
            )
        )

    source_sha256 = sha256_file(path)
    contract = FinancialCalendarContract(
        source_sha256=source_sha256,
        source_path=str(path),
        version=f"financial-calendar.v2:{source_sha256[:12]}",
        company_id=session_company_id,
        events=events,
        links=links,
    )
    _validate_contract(contract)
    return contract
