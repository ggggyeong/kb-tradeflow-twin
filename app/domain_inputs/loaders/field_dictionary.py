from __future__ import annotations

import hashlib
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.schemas.field_contract import (
    PROJECTION_TARGETS,
    DocumentType,
    FieldRegistry,
    FieldRegistryEntry,
    to_safe_key,
)

MAIN_SHEET = "BL, Booking, Invoice"
CORE_SHEET = "필수 - 없으면 안됨"
EXPECTED_MAIN_COUNTS = {
    DocumentType.BILL_OF_LADING: 23,
    DocumentType.BOOKING_CONFIRMATION: 52,
    DocumentType.COMMERCIAL_INVOICE: 23,
}
EXPECTED_CORE_COUNTS = {
    DocumentType.BILL_OF_LADING: 8,
    DocumentType.BOOKING_CONFIRMATION: 8,
    DocumentType.COMMERCIAL_INVOICE: 6,
}
EXPECTED_HEADERS = [
    "required",
    "doc_type",
    "standard_field",
    "aliases",
    "meaning",
    "Example",
]


def sha256_file(path: Path) -> str:
    """Calculate a binary file SHA-256 without altering the source."""
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _rows(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Missing field dictionary sheet: {sheet_name}")
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    if headers[:6] != EXPECTED_HEADERS:
        raise ValueError(f"Invalid headers on {sheet_name}: {headers[:6]}")

    current_required: str | None = None
    current_doc_type: DocumentType | None = None
    output: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        required, doc_type, standard_field, aliases, meaning, example = values[:6]
        required_text = str(required).strip() if required is not None else ""
        if required_text in {"Y", "N"}:
            current_required = required_text
        doc_text = str(doc_type).strip() if doc_type is not None else ""
        if doc_text in {item.value for item in DocumentType}:
            current_doc_type = DocumentType(doc_text)
        if standard_field is None or not str(standard_field).strip():
            continue
        if current_required not in {"Y", "N"} or current_doc_type is None:
            raise ValueError(f"Unscoped field row on {sheet_name}: {values[:3]}")
        output.append(
            {
                "required": current_required,
                "doc_type": current_doc_type,
                "standard_field": str(standard_field),
                "aliases": str(aliases or ""),
                "meaning": str(meaning or ""),
                "example": str(example or ""),
            }
        )
    return output


def load_field_registry(path: Path) -> FieldRegistry:
    """Load and validate the 98 main rows and 22 actual core rows."""
    main_rows = _rows(path, MAIN_SHEET)
    core_rows = _rows(path, CORE_SHEET)
    main_counts = Counter(row["doc_type"] for row in main_rows)
    core_counts = Counter(row["doc_type"] for row in core_rows)
    if len(main_rows) != 98 or main_counts != Counter(EXPECTED_MAIN_COUNTS):
        raise ValueError(f"Field main count mismatch: total={len(main_rows)}, {main_counts}")
    if len(core_rows) != 22 or core_counts != Counter(EXPECTED_CORE_COUNTS):
        raise ValueError(f"Core field count mismatch: total={len(core_rows)}, {core_counts}")

    core_keys = {(row["doc_type"], row["standard_field"]) for row in core_rows}
    main_keys = {(row["doc_type"], row["standard_field"]) for row in main_rows}
    if not core_keys <= main_keys:
        raise ValueError(f"Core keys missing from main sheet: {sorted(core_keys - main_keys)}")
    if len(main_keys) != 98:
        raise ValueError("Duplicate (doc_type, exact_standard_field) in main sheet")

    safe_by_doc: dict[DocumentType, set[str]] = {doc_type: set() for doc_type in DocumentType}
    entries: list[FieldRegistryEntry] = []
    for row in main_rows:
        doc_type = row["doc_type"]
        exact = row["standard_field"]
        safe_key = to_safe_key(exact)
        if safe_key in safe_by_doc[doc_type]:
            raise ValueError(f"Safe key collision: {doc_type.value}::{safe_key}")
        safe_by_doc[doc_type].add(safe_key)
        aliases = [part.strip() for part in row["aliases"].split(" / ") if part.strip()]
        entries.append(
            FieldRegistryEntry(
                doc_type=doc_type,
                exact_standard_field=exact,
                safe_key=safe_key,
                required_in_document=row["required"] == "Y",
                core_blocking_candidate=(doc_type, exact) in core_keys,
                aliases=aliases,
                meaning=row["meaning"],
                example=row["example"],
                projection_target=PROJECTION_TARGETS.get((doc_type, exact)),
            )
        )

    registry = FieldRegistry(
        entries=entries,
        source_sha256=sha256_file(path),
        source_path=str(path),
    )
    protected = {
        entry.exact_standard_field: entry.safe_key
        for entry in registry.for_document(DocumentType.COMMERCIAL_INVOICE)
        if entry.exact_standard_field in {"invoice No.", "shipping Marks", "signed by"}
    }
    if protected != {
        "invoice No.": "invoice_no",
        "shipping Marks": "shipping_marks",
        "signed by": "signed_by",
    }:
        raise ValueError(f"Protected exact keys changed: {protected}")
    return registry
