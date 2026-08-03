from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.domain_inputs.loaders.field_dictionary import sha256_file
from app.schemas.financial_calendar import (
    DayBand,
    FinancialPriorityPolicy,
    PriorityExampleRow,
    ResponsePriorityBand,
)

EXPECTED_SHEETS = [
    "1.우선순위 로직 스펙",
    "2.긴급도 구간 정의(표시 전용)",
    "3.이벤트 유형별 기본 영향도",
    "4.impact_level 산출 규칙",
    "5.response_priority_level 등급",
    "6.적용 예시(재정렬)",
    "7.KB 라우팅 규칙",
]
EVENT_TYPE_CODES = {
    "운전자금대출 만기": "WORKING_CAPITAL_LOAN_MATURITY",
    "선물환 계약 만기": "FX_FORWARD_MATURITY",
    "공급업체 지급일": "SUPPLIER_PAYMENT",
}


def _rows(sheet: Worksheet) -> list[dict[str, Any]]:
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        raise ValueError(f"Required priority sheet is empty: {sheet.title}")
    headers = [str(value or "").strip() for value in values[0]]
    if not headers[0]:
        raise ValueError(f"Priority sheet has no first header: {sheet.title}")
    return [
        dict(zip(headers, row, strict=True))
        for row in values[1:]
        if any(value not in (None, "") for value in row)
    ]


def _range_for_label(label: str) -> tuple[int | None, int | None]:
    normalized = label.upper().replace(" ", "")
    if "OVERDUE" in normalized and "D-0" in normalized:
        return None, 0
    if normalized == "OVERDUE":
        return None, -1
    if normalized in {"D-0", "D0"}:
        return 0, 0
    match = re.fullmatch(r"D(\d+)-(\d+)", normalized)
    if match:
        return int(match.group(1)), int(match.group(2))
    match = re.fullmatch(r"D(\d+)\+", normalized)
    if match:
        return int(match.group(1)), None
    raise ValueError(f"Unsupported days-until-event band: {label!r}")


def _event_type_code(label: str) -> str:
    for event_name, code in EVENT_TYPE_CODES.items():
        if label.startswith(event_name):
            return code
    raise ValueError(f"Unsupported event type label: {label}")


def _routing_key(event_type_code: str, is_kb_contract: Any) -> str:
    if is_kb_contract is None:
        suffix = "ANY"
    elif isinstance(is_kb_contract, bool):
        suffix = "TRUE" if is_kb_contract else "FALSE"
    else:
        raise ValueError(
            f"is_kb_contract in priority routing must be Boolean or blank, got {is_kb_contract!r}"
        )
    return f"{event_type_code}:{suffix}"


def load_financial_priority_policy(path: Path) -> FinancialPriorityPolicy:
    """Load the deterministic v3 ordering and routing policy from its workbook."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    if workbook.sheetnames != EXPECTED_SHEETS:
        raise ValueError(f"Financial priority sheet mismatch: {workbook.sheetnames}")

    spec = workbook[EXPECTED_SHEETS[0]]
    if (
        spec["B3"].value != "days_until_event — 정확한 일수"
        or "CONFIRMED > UNCONFIRMED" not in str(spec["C4"].value)
        or "event_id 오름차순" not in str(spec["B8"].value)
    ):
        raise ValueError("Priority sort contract does not match reviewed v3 semantics")

    urgency_rows = _rows(workbook[EXPECTED_SHEETS[1]])
    urgency_bands: list[DayBand] = []
    for row in urgency_rows:
        label = str(row["구간 코드"]).strip()
        if not label or label.startswith("["):
            continue
        min_days, max_days = _range_for_label(label)
        urgency_bands.append(DayBand(label=label, min_days=min_days, max_days=max_days))

    type_rows = _rows(workbook[EXPECTED_SHEETS[2]])
    event_type_order: dict[str, int] = {}
    for row in type_rows:
        event_name = str(row["이벤트 유형"]).strip()
        if not event_name or event_name.startswith("["):
            continue
        order_match = re.match(r"(\d+)순위", str(row["기본순서"]).strip())
        if order_match is None:
            raise ValueError(f"Invalid event type order: {row['기본순서']!r}")
        event_type_order[_event_type_code(event_name)] = int(order_match.group(1))

    impact_sheet = workbook[EXPECTED_SHEETS[3]]
    expected_impact_rules = {
        "C10": "UNKNOWN",
        "C11": "CRITICAL",
        "C12": "HIGH",
        "C13": "MEDIUM",
        "C14": "LOW",
    }
    if any(
        str(impact_sheet[cell].value).strip() != expected
        for cell, expected in expected_impact_rules.items()
    ) or not str(impact_sheet["C15"].value).startswith("REVIEW_REQUIRED"):
        raise ValueError("Impact-level rules differ from the reviewed v3 contract")

    priority_rows = _rows(workbook[EXPECTED_SHEETS[4]])
    response_bands: list[ResponsePriorityBand] = []
    for row in priority_rows:
        label = str(row["days_until_event 범위"]).strip()
        priority = row["response_priority_level"]
        if not label or label.startswith("[") or priority not in {"P1", "P2", "P3", "P4"}:
            continue
        normalized_label = "OVERDUE 또는 D-0" if label.startswith("OVERDUE") else label
        min_days, max_days = _range_for_label(normalized_label)
        response_bands.append(
            ResponsePriorityBand(
                label=normalized_label,
                min_days=min_days,
                max_days=max_days,
                priority=priority,
            )
        )

    example_rows = _rows(workbook[EXPECTED_SHEETS[5]])
    portfolio_examples: list[PriorityExampleRow] = []
    for row in example_rows:
        rank = row["expected_priority_rank"]
        if not isinstance(rank, int) or row["ranking_scope"] != "KB_PORTFOLIO_VIEW":
            continue
        portfolio_examples.append(
            PriorityExampleRow(
                expected_priority_rank=rank,
                expected_response_priority_level=str(row["expected_response_priority_level"]),
                company_id=str(row["company_id"]),
                event_id=str(row["event_id"]),
                days_until_event=int(row["days_until_event"]),
                urgency_bucket=str(row["urgency_bucket(표시용)"]),
                link_status=str(row["link_status"]),
                impact_level=str(row["impact_level"]),
                event_type=str(row["이벤트 유형"]),
                lead_days=int(row["lead_days"]),
                same_day_flag=row["SAME_DAY"] == "예",
                ranking_scope=str(row["ranking_scope"]),
                display_label=str(row["표시 라벨"]),
            )
        )

    routing_rows = _rows(workbook[EXPECTED_SHEETS[6]])
    routing: dict[str, dict[str, str]] = {}
    for row in routing_rows:
        code = str(row["event_type_code"]).strip()
        if not code or code.startswith("["):
            continue
        if code not in EVENT_TYPE_CODES.values():
            if row.get("action_owner") in (None, "") and row.get("handoff_message") in (None, ""):
                continue
            raise ValueError(f"Unsupported routing event_type_code: {code}")
        routing[_routing_key(code, row["is_kb_contract"])] = {
            "action_owner": str(row["action_owner"]).strip(),
            "handoff_message": str(row["handoff_message"]).strip(),
        }

    source_hash = sha256_file(path)
    return FinancialPriorityPolicy(
        source_sha256=source_hash,
        source_path=str(path),
        version=f"financial-priority.v3:{source_hash[:12]}",
        response_bands=response_bands,
        urgency_bands=urgency_bands,
        event_type_order=event_type_order,
        link_status_order={"CONFIRMED": 0, "UNCONFIRMED": 1},
        impact_level_order={
            "CRITICAL": 0,
            "HIGH": 1,
            "MEDIUM": 2,
            "LOW": 3,
            "UNKNOWN": 4,
            "REVIEW_REQUIRED": 5,
        },
        routing=routing,
        portfolio_examples=portfolio_examples,
    )
