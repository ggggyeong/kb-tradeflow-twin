from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.db.models import (
    Base,
    Company,
    FinancialEvent,
    FinancialTransactionTimeline,
    PaymentObligation,
    TradeCase,
    TransactionFinancialEventLink,
)
from app.domain_inputs.loaders.financial_calendar import (
    EVENT_HEADERS,
    LINK_HEADERS,
    REQUIRED_SHEETS,
    load_financial_calendar_contract,
)
from app.services.financial_calendar import FinancialCalendarService
from app.services.ingestion import batch_service


def _write_two_sheet_workbook(
    path: Path,
    *,
    transaction_id: str = "TXN-A-001",
    event_id: str = "EVENT-A-001",
) -> Path:
    workbook = Workbook()
    event_sheet = workbook.active
    event_sheet.title = REQUIRED_SHEETS[0]
    event_sheet.append(EVENT_HEADERS)
    event_sheet.append(
        [
            event_id,
            "WORKING_CAPITAL_LOAN_MATURITY",
            date(2026, 8, 20),
            50_000,
            "USD",
            "KB국민은행",
        ]
    )

    link_sheet = workbook.create_sheet(REQUIRED_SHEETS[1])
    link_sheet.append(LINK_HEADERS)
    link_sheet.append(
        [
            transaction_id,
            event_id,
            "EXPECTED_EXPORT_RECEIPT",
            "CONFIRMED",
        ]
    )
    workbook.save(path)
    return path


def _seed_case(
    session: Session,
    *,
    company_id: str,
    transaction_id: str,
    case_id: str,
    monitoring_enabled: bool = True,
) -> TradeCase:
    company = session.get(Company, company_id)
    if company is None:
        company = Company(company_id=company_id, legal_name=f"{company_id} 주식회사")
        session.add(company)
        session.flush()
    trade_case = TradeCase(
        case_id=case_id,
        company_id=company_id,
        transaction_id=transaction_id,
        company=company.legal_name,
        counterparty="Overseas Buyer",
        invoice_no=f"INV-{transaction_id}",
        currency="USD",
        status="ACTIVE",
        monitoring_enabled=monitoring_enabled,
        basis_version="basis-test-1",
    )
    session.add(trade_case)
    session.flush()
    return trade_case


def _memory_session() -> Session:
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    return Session(engine)


def test_uploaded_financial_calendar_batch_requires_one_xlsx(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    upload_root = tmp_path / "uploads"
    monkeypatch.setattr(batch_service, "UPLOAD_ROOT", upload_root)
    target = batch_service.upload_batch_directory("FINANCIAL-CALENDAR-BATCH")
    target.mkdir(parents=True)
    workbook = target / "financial-calendar.xlsx"
    workbook.write_bytes(b"xlsx")

    assert (
        FinancialCalendarService.resolve_workbook_path(batch_id="FINANCIAL-CALENDAR-BATCH")
        == workbook
    )

    (target / "booking.pdf").write_bytes(b"pdf")
    with pytest.raises(ValueError, match="exactly one XLSX"):
        FinancialCalendarService.resolve_workbook_path(batch_id="FINANCIAL-CALENDAR-BATCH")


@pytest.mark.parametrize("invalid_layout", ["extra_sheet", "wrong_header", "reversed_order"])
def test_two_sheet_contract_requires_exact_sheets_and_headers(
    tmp_path: Path,
    invalid_layout: str,
) -> None:
    path = _write_two_sheet_workbook(tmp_path / f"{invalid_layout}.xlsx")
    # Reloading through openpyxl keeps the test independent from workbook XML details.
    workbook = load_workbook(path)
    if invalid_layout == "extra_sheet":
        workbook.create_sheet("3.허용되지않는시트")
    elif invalid_layout == "wrong_header":
        workbook[REQUIRED_SHEETS[0]].cell(row=1, column=1, value="company_id")
    else:
        workbook.move_sheet(workbook[REQUIRED_SHEETS[1]], offset=-1)
    workbook.save(path)

    with pytest.raises(ValueError, match=r"exactly|Header mismatch"):
        load_financial_calendar_contract(path, company_id="COMP-A")


def test_two_sheet_loader_derives_company_and_server_owned_fields(tmp_path: Path) -> None:
    path = _write_two_sheet_workbook(tmp_path / "financial-calendar.xlsx")

    contract = load_financial_calendar_contract(path, company_id="COMP-A")

    assert contract.company_id == "COMP-A"
    assert contract.version.startswith("financial-calendar.v2:")
    assert contract.events[0].event_name == "운전자금대출 만기"
    assert contract.events[0].is_kb_contract is True
    assert contract.links[0].link_id.startswith("LNK-")
    assert contract.links[0].dependency_scope == "UNKNOWN"


def test_validation_checks_database_ownership_without_writing(tmp_path: Path) -> None:
    path = _write_two_sheet_workbook(tmp_path / "financial-calendar.xlsx")
    with _memory_session() as session:
        _seed_case(
            session,
            company_id="COMP-A",
            transaction_id="TXN-A-001",
            case_id="CASE-A-001",
        )
        service = FinancialCalendarService(session)

        validation = service.validate_workbook(path, company_id="COMP-A")

        assert validation["valid"] is True
        assert validation["company_id"] == "COMP-A"
        assert validation["transaction_count"] == 1
        assert validation["event_count"] == 1
        assert validation["link_count"] == 1
        assert session.scalar(select(func.count()).select_from(FinancialEvent)) == 0
        assert session.scalar(select(func.count()).select_from(TransactionFinancialEventLink)) == 0


def test_import_writes_only_events_and_links_and_is_idempotent(tmp_path: Path) -> None:
    path = _write_two_sheet_workbook(tmp_path / "financial-calendar.xlsx")
    with _memory_session() as session:
        trade_case = _seed_case(
            session,
            company_id="COMP-A",
            transaction_id="TXN-A-001",
            case_id="CASE-A-001",
        )
        service = FinancialCalendarService(session)

        first = service.import_workbook(path, company_id="COMP-A")
        second = service.import_workbook(path, company_id="COMP-A")

        assert first["created"]["events"] == 1
        assert first["created"]["links"] == 1
        assert second["created"]["events"] == 0
        assert second["created"]["links"] == 0
        assert session.scalar(select(func.count()).select_from(Company)) == 1
        assert session.scalar(select(func.count()).select_from(TradeCase)) == 1
        assert session.scalar(select(func.count()).select_from(FinancialEvent)) == 1
        assert session.scalar(select(func.count()).select_from(TransactionFinancialEventLink)) == 1
        link = session.scalar(select(TransactionFinancialEventLink))
        assert link is not None
        assert (link.company_id, link.case_id, link.transaction_id) == (
            "COMP-A",
            "CASE-A-001",
            "TXN-A-001",
        )
        assert link.dependency_scope == "UNKNOWN"
        assert link.linked_amount is None
        assert link.linked_currency is None
        assert session.scalar(select(func.count()).select_from(FinancialTransactionTimeline)) == 0
        assert session.scalar(select(func.count()).select_from(PaymentObligation)) == 0
        session.refresh(trade_case)
        assert trade_case.case_id == "CASE-A-001"
        assert trade_case.status == "ACTIVE"


def test_post_conflict_human_review_updates_optional_link_facts_and_survives_reimport(
    tmp_path: Path,
) -> None:
    path = _write_two_sheet_workbook(tmp_path / "financial-calendar.xlsx")
    with _memory_session() as session:
        _seed_case(
            session,
            company_id="COMP-A",
            transaction_id="TXN-A-001",
            case_id="CASE-A-001",
        )
        service = FinancialCalendarService(session)
        service.import_workbook(path, company_id="COMP-A")
        link = session.scalar(select(TransactionFinancialEventLink))
        assert link is not None

        scoped = service.update_financial_event_link_details(
            company_id="COMP-A",
            case_id="CASE-A-001",
            transaction_id="TXN-A-001",
            transaction_event_link_id=link.transaction_event_link_id,
            details={"dependency_scope": "PARTIAL"},
            actor="human:test-thread",
        )
        assert scoped["complete"] is False
        assert scoped["missing_fields"] == ["linked_amount", "linked_currency"]

        completed = service.update_financial_event_link_details(
            company_id="COMP-A",
            case_id="CASE-A-001",
            transaction_id="TXN-A-001",
            transaction_event_link_id=link.transaction_event_link_id,
            details={"linked_amount": 30_000, "linked_currency": "usd"},
            actor="human:test-thread",
        )
        assert completed["complete"] is True
        assert completed["linked_amount"] == 30_000
        assert completed["linked_currency"] == "USD"

        service.import_workbook(path, company_id="COMP-A")
        session.refresh(link)
        assert link.dependency_scope == "PARTIAL"
        assert link.linked_amount == 30_000
        assert link.linked_currency == "USD"
        assert link.dependency_basis == "HUMAN_CONFIRMED"
        assert link.confirmed_by == "human:test-thread"
        assert link.source_metadata_json["human_link_review"]["reviewed_fields"] == [
            "linked_amount",
            "linked_currency",
        ]

        with pytest.raises(ValueError, match="ownership mismatch"):
            service.update_financial_event_link_details(
                company_id="COMP-B",
                case_id="CASE-A-001",
                transaction_id="TXN-A-001",
                transaction_event_link_id=link.transaction_event_link_id,
                details={"dependency_scope": "FULL"},
                actor="human:test-thread",
            )


def test_validation_rejects_cross_company_transaction(tmp_path: Path) -> None:
    path = _write_two_sheet_workbook(
        tmp_path / "cross-company.xlsx",
        transaction_id="TXN-B-001",
    )
    with _memory_session() as session:
        _seed_case(
            session,
            company_id="COMP-A",
            transaction_id="TXN-A-001",
            case_id="CASE-A-001",
        )
        _seed_case(
            session,
            company_id="COMP-B",
            transaction_id="TXN-B-001",
            case_id="CASE-B-001",
        )

        with pytest.raises(ValueError, match=r"[Cc]ross-company|different company"):
            FinancialCalendarService(session).validate_workbook(path, company_id="COMP-A")


def test_validation_rejects_unknown_transaction(tmp_path: Path) -> None:
    path = _write_two_sheet_workbook(
        tmp_path / "unknown-transaction.xlsx",
        transaction_id="TXN-UNKNOWN",
    )
    with _memory_session() as session:
        _seed_case(
            session,
            company_id="COMP-A",
            transaction_id="TXN-A-001",
            case_id="CASE-A-001",
        )

        with pytest.raises(ValueError, match=r"[Uu]nknown transaction"):
            FinancialCalendarService(session).validate_workbook(path, company_id="COMP-A")


def test_candidate_selection_is_isolated_by_session_company(tmp_path: Path) -> None:
    company_a_path = _write_two_sheet_workbook(
        tmp_path / "company-a.xlsx",
        transaction_id="TXN-A-001",
        event_id="EVENT-A-001",
    )
    company_b_path = _write_two_sheet_workbook(
        tmp_path / "company-b.xlsx",
        transaction_id="TXN-B-001",
        event_id="EVENT-B-001",
    )
    with _memory_session() as session:
        _seed_case(
            session,
            company_id="COMP-A",
            transaction_id="TXN-A-001",
            case_id="CASE-A-001",
        )
        _seed_case(
            session,
            company_id="COMP-B",
            transaction_id="TXN-B-001",
            case_id="CASE-B-001",
        )
        service = FinancialCalendarService(session)
        service.import_workbook(company_a_path, company_id="COMP-A")
        service.import_workbook(company_b_path, company_id="COMP-B")

        selected = service.select_monitoring_candidates(
            company_id="COMP-A",
            as_of_date=date(2026, 8, 1),
        )

        assert selected["candidate_count"] == 1
        assert selected["candidates"] == [
            {
                "case_id": "CASE-A-001",
                "company_id": "COMP-A",
                "transaction_id": "TXN-A-001",
                "basis_version": "basis-test-1",
            }
        ]
