from pathlib import Path
from typing import Any

import pytest

from app.schemas.portfolio import PortfolioRunRequest
from app.services import financial_calendar as module

ROOT = Path(__file__).resolve().parents[2]
EXAMPLE = ROOT / "data/fixtures/tradeflow_example/financial_calendar.xlsx"


def test_example_reads_typed_dates_and_only_selected_transaction() -> None:
    calendar = module.read_financial_calendar(EXAMPLE, "DEMO-EXPORT-001")
    assert [e.event_id for e in calendar.events] == ["DEMO-LOAN", "DEMO-FX", "DEMO-PAY"]
    assert str(calendar.events[0].event_date) == "2026-10-15"
    assert calendar.events[0].amount == 30_000_000
    assert calendar.events[1].fx_direction == "SELL"
    assert calendar.events[2].payment_purpose == "DOMESTIC"
    assert calendar.document_references["invoice_no"] == "INV-DEMO-001"
    assert not calendar.warnings


@pytest.mark.parametrize(
    ("folder", "trade_id"),
    [("demo_1_manual_today_risk", "TXN-DEMO1"), ("demo_2_reported_delay", "TXN-DEMO2")],
)
def test_original_two_sheet_calendar_still_works(folder: str, trade_id: str) -> None:
    old = ROOT / "data/judge_demo_final" / folder / "financial_calendar.xlsx"
    calendar = module.read_financial_calendar(old, trade_id)
    assert len(calendar.events) == 1
    assert calendar.events[0].link_status == "CONFIRMED"
    assert not calendar.document_references and calendar.warnings


def test_unknown_transaction_is_not_guessed() -> None:
    with pytest.raises(ValueError, match="연결된 금융일정"):
        module.read_financial_calendar(EXAMPLE, "DOES-NOT-EXIST")


@pytest.mark.parametrize(
    "case",
    [
        "duplicate_event",
        "missing_event",
        "duplicate_link",
        "wrong_scenario",
        "bad_status",
        "wrong_link_type",
        "ambiguous_date",
        "duplicate_trade",
        "too_many_events",
    ],
)
def test_bad_calendar_records_fail_closed(monkeypatch: Any, case: str) -> None:
    original = module._rows

    def changed(workbook: Any, sheet: str, required: set[str]) -> list[dict[str, Any]]:
        rows = original(workbook, sheet, required)
        if sheet == module.EVENT_SHEET:
            if case == "duplicate_event":
                rows.append(dict(rows[0]))
            if case == "wrong_scenario":
                rows[0]["event_type_code"] = "INVENTED"
            if case == "ambiguous_date":
                rows[0]["event_date"] = "10/11/2026"
        if sheet == module.LINK_SHEET:
            if case == "missing_event":
                rows[0]["event_id"] = "MISSING"
            if case == "duplicate_link":
                rows.append(dict(rows[0]))
            if case == "bad_status":
                rows[0]["link_status"] = "ASSUME_CONFIRMED"
            if case == "wrong_link_type":
                rows[0]["link_type"] = "FX_SETTLEMENT_SOURCE"
            if case == "too_many_events":
                rows[-1]["transaction_id"] = rows[0]["transaction_id"]
        if sheet == module.TRADE_SHEET and case == "duplicate_trade":
            rows.append(dict(rows[0]))
        return rows

    monkeypatch.setattr(module, "_rows", changed)
    with pytest.raises(ValueError):
        module.read_financial_calendar(EXAMPLE, "DEMO-EXPORT-001")


@pytest.mark.parametrize(
    "changes",
    [
        {},
        {
            "transaction_id": "DEMO",
            "financial_events": [
                {"event_id": "x", "event_name": "x", "scenario_code": "SUPPLIER_PAYMENT"}
            ],
        },
    ],
)
def test_excel_requires_trade_id_and_rejects_mixed_schedule_sources(
    changes: dict[str, Any],
) -> None:
    with pytest.raises(ValueError):
        PortfolioRunRequest(financial_calendar_path=EXAMPLE, **changes)


def test_formula_input_is_rejected_without_using_cached_values(monkeypatch: Any) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(EXAMPLE)
    workbook[module.EVENT_SHEET]["C2"] = "=TODAY()"
    monkeypatch.setattr("openpyxl.load_workbook", lambda *args, **kwargs: workbook)
    with pytest.raises(ValueError, match="수식"):
        module.read_financial_calendar(EXAMPLE, "DEMO-EXPORT-001")


def test_missing_required_header_is_rejected(monkeypatch: Any) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(EXAMPLE)
    workbook[module.LINK_SHEET]["D1"] = "unknown_header"
    monkeypatch.setattr("openpyxl.load_workbook", lambda *args, **kwargs: workbook)
    with pytest.raises(ValueError, match="필수 열"):
        module.read_financial_calendar(EXAMPLE, "DEMO-EXPORT-001")


def test_sparse_oversize_workbook_is_rejected(monkeypatch: Any) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(EXAMPLE)
    workbook[module.LINK_SHEET]["A2000"] = "HIDDEN-TRANSACTION"
    monkeypatch.setattr("openpyxl.load_workbook", lambda *args, **kwargs: workbook)
    with pytest.raises(ValueError, match="1,000"):
        module.read_financial_calendar(EXAMPLE, "DEMO-EXPORT-001")
